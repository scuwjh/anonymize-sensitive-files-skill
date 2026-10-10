#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from pathlib import Path
from typing import Iterable

from detectors import (
    FieldRule,
    Finding,
    detect_field_rule_values,
    detect_sensitive_values,
    field_rule_matches_label,
    load_custom_terms,
    load_field_rules,
)
from fake_data import KNOWN_FAKE_VALUES, FakeDataFactory, normalize_category


SUPPORTED_TEXT_EXTENSIONS = {".md", ".markdown", ".txt", ".text", ".csv", ".tsv", ".json", ".yaml", ".yml"}
SUPPORTED_EXCEL_EXTENSIONS = {".xlsx", ".xlsm"}
SUPPORTED_EXTENSIONS = SUPPORTED_TEXT_EXTENSIONS | {".docx", ".pdf"} | SUPPORTED_EXCEL_EXTENSIONS
DOCX_METADATA_FIELDS = (
    "author",
    "last_modified_by",
    "title",
    "subject",
    "comments",
    "keywords",
    "category",
)
DOCX_PART_NAMES = (
    "header",
    "footer",
    "first_page_header",
    "first_page_footer",
    "even_page_header",
    "even_page_footer",
)
STRUCTURED_CATEGORIES = frozenset(
    {"phone", "email", "id_card", "passport", "bank_card", "bank_routing_code", "social_credit_code", "ip", "url", "domain", "secret"}
)
FAKE_VALUE_PATTERN = re.compile(
    "|".join(sorted((re.escape(value) for value in KNOWN_FAKE_VALUES if len(value) >= 2), key=len, reverse=True))
)

PDF_METADATA_FIELDS = ("title", "author", "subject", "keywords", "creator")
EXCEL_HYPERLINK_TEXT_ATTRIBUTES = ("target", "location", "display", "tooltip")
EXCEL_WORKBOOK_PROPERTY_ATTRIBUTES = (
    "creator",
    "lastModifiedBy",
    "title",
    "subject",
    "description",
    "keywords",
    "category",
    "contentStatus",
    "identifier",
    "language",
    "version",
    "revision",
)
EXCEL_TABLE_HEADER_VALUES = frozenset({"项目", "日期", "签字", "签名", "盖章", "签章", "字段", "值"})
INVALID_EXCEL_SHEET_TITLE_CHARS = re.compile(r"[\[\]:*?/\\]")


class ProcessingError(RuntimeError):
    def __init__(self, message: str, findings: list[Finding] | None = None, warnings: list[str] | None = None):
        super().__init__(message)
        self.findings = findings or []
        self.warnings = warnings or []


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    if args.export_mapping and not args.mapping:
        print(
            json.dumps({"error": "--export-mapping requires --mapping so raw values are not printed to stdout."}, ensure_ascii=False),
            file=sys.stderr,
        )
        return 2
    if args.no_field_rules and args.field_rules:
        print(json.dumps({"error": "--no-field-rules cannot be combined with --field-rules."}, ensure_ascii=False), file=sys.stderr)
        return 2

    inputs, skipped_inputs = collect_inputs(args.inputs, args.recursive)
    if not inputs:
        print(
            json.dumps(
                {
                    "error": "No supported input files found.",
                    "summary": skipped_summary(skipped_inputs),
                    "skipped_inputs": skipped_inputs,
                },
                ensure_ascii=False,
                indent=2,
            ),
            file=sys.stderr,
        )
        return 2
    try:
        custom_terms = load_custom_terms(args.terms)
        field_rules = load_field_rules(args.field_rules, use_default=not args.no_field_rules)
    except Exception as exc:
        print(
            json.dumps({"error": safe_error_message(str(exc), *argument_paths(args))}, ensure_ascii=False, indent=2),
            file=sys.stderr,
        )
        return 2
    factory = FakeDataFactory(seed=args.seed)
    factory.reserve(value for _, value in custom_terms)

    try:
        output_paths = plan_outputs(inputs, args, custom_terms, field_rules, factory)
    except ValueError as exc:
        print(json.dumps({"error": str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 2
    results = []

    for path in inputs:
        output_path = output_paths.get(path)
        try:
            results.append(process_file(path, args, custom_terms, field_rules, factory, output_path))
        except ProcessingError as exc:
            result = result_for(path, None, "error", exc.findings, exc.warnings)
            result["error"] = safe_error_message(str(exc), *paths_for_error(args, path, output_path))
            results.append(result)
        except Exception as exc:
            result = result_for(path, None, "error", [], [])
            result["error"] = safe_error_message(str(exc), *paths_for_error(args, path, output_path))
            results.append(result)

    report = build_report(results, factory, skipped_inputs)

    try:
        if args.report:
            args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        if args.mapping:
            mapping = factory.sensitive_mapping() if args.export_mapping else factory.safe_mapping()
            args.mapping.write_text(json.dumps(mapping, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception as exc:
        print(
            json.dumps({"error": safe_error_message(str(exc), *argument_paths(args))}, ensure_ascii=False, indent=2),
            file=sys.stderr,
        )
        return 1

    print(json.dumps(report, ensure_ascii=False, indent=2))
    return exit_code_for(results, skipped_inputs)


def parse_args(argv: list[str] | None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Scan or anonymize sensitive data in text, Markdown, DOCX, PDF, and Excel files.",
    )
    parser.add_argument("inputs", nargs="+", type=Path, help="File or directory paths to process.")
    parser.add_argument(
        "--mode",
        choices=["scan", "anonymize", "verify"],
        default="anonymize",
        help="scan reports findings only; anonymize writes output files; verify scans files after anonymization.",
    )
    parser.add_argument("--recursive", action="store_true", help="Recurse into input directories.")
    parser.add_argument("--output", type=Path, help="Output file path for a single input file.")
    parser.add_argument("--output-dir", type=Path, help="Directory for anonymized output files.")
    parser.add_argument("--report", type=Path, help="Write the JSON report to this path.")
    parser.add_argument("--terms", type=Path, help="Custom terms file. Lines may be 'category:value' or plain values.")
    parser.add_argument(
        "--field-rules",
        type=Path,
        help="YAML file containing global label/value field rules. Defaults to the bundled rules.",
    )
    parser.add_argument("--no-field-rules", action="store_true", help="Disable bundled and explicit field rules.")
    parser.add_argument(
        "--anonymize-filenames",
        action="store_true",
        help="Anonymize output file names with the same mapping so file names cannot leak entities.",
    )
    parser.add_argument("--seed", help="Optional seed label used for stable safe mapping hashes.")
    parser.add_argument("--mapping", type=Path, help="Optional mapping output path. Required with --export-mapping.")
    parser.add_argument(
        "--export-mapping",
        action="store_true",
        help="Write raw original-to-fake mappings to --mapping. This creates a sensitive artifact.",
    )
    return parser.parse_args(argv)


def argument_paths(args: argparse.Namespace) -> list[Path]:
    paths = list(args.inputs)
    for attribute in ("output", "output_dir", "report", "terms", "field_rules", "mapping"):
        value = getattr(args, attribute, None)
        if value is not None:
            paths.append(value)
    return paths


def paths_for_error(args: argparse.Namespace, input_path: Path, output_path: Path | None = None) -> list[Path]:
    paths = [input_path, *argument_paths(args)]
    if output_path is not None:
        paths.append(output_path)
    return paths


def resolve_inputs(inputs: list[Path], recursive: bool) -> Iterable[Path]:
    supported, _skipped = collect_inputs(inputs, recursive)
    yield from supported


def collect_inputs(inputs: list[Path], recursive: bool) -> tuple[list[Path], list[dict]]:
    supported: list[Path] = []
    skipped: list[dict] = []
    # Fake values are assigned in first-appearance order, and output collisions are
    # resolved in traversal order, so the input sequence is part of the reproducibility
    # contract. Directory walks follow filesystem order, which changes when a corpus is
    # copied or re-extracted; normalize every collected path by sorted string order.
    unordered: list[Path] = []
    for item in inputs:
        item = item.expanduser()
        if item.is_dir():
            iterator = item.rglob("*") if recursive else item.glob("*")
            for child in iterator:
                if not child.is_file():
                    continue
                if child.suffix.lower() in SUPPORTED_EXTENSIONS:
                    unordered.append(child)
                else:
                    skipped.append(skipped_input(child, "unsupported_extension", explicit=False))
        elif item.is_file():
            if item.suffix.lower() in SUPPORTED_EXTENSIONS:
                unordered.append(item)
            else:
                skipped.append(skipped_input(item, "unsupported_extension", explicit=True))
        else:
            skipped.append(skipped_input(item, "missing", explicit=True))
    seen_paths: set[str] = set()
    for path in sorted(unordered, key=str):
        key = str(path.expanduser())
        if key in seen_paths:
            continue
        seen_paths.add(key)
        supported.append(path)
    return supported, skipped


def skipped_input(path: Path, reason: str, explicit: bool) -> dict:
    item = {"path": safe_report_path(path), "reason": reason, "explicit": explicit}
    if path.suffix:
        item["extension"] = path.suffix.lower()
    return item


def process_file(
    path: Path,
    args: argparse.Namespace,
    custom_terms: list[tuple[str, str]],
    field_rules: list[FieldRule],
    factory: FakeDataFactory,
    output_path: Path | None,
) -> dict:
    suffix = path.suffix.lower()
    if args.mode == "scan":
        findings, warnings = scan_file(path, custom_terms, field_rules)
        return result_for(path, None, "scanned", findings, warnings)

    if args.mode == "verify":
        findings, warnings = verify_file(path, custom_terms, field_rules)
        status = "verified" if not findings else "residual_findings"
        return result_for(path, None, status, findings, warnings)

    if output_path is None:
        raise ValueError("Missing output path for anonymization.")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if suffix in SUPPORTED_TEXT_EXTENSIONS:
        text = path.read_text(encoding="utf-8")
        anonymized, findings = anonymize_text(text, custom_terms, field_rules, factory)
        output_path.write_text(anonymized, encoding="utf-8")
        return result_for(path, output_path, "anonymized", findings, [])

    if suffix == ".docx":
        findings, warnings = anonymize_docx(path, output_path, custom_terms, field_rules, factory)
        return result_for(path, output_path, "anonymized", findings, warnings)

    if suffix == ".pdf":
        findings, warnings = anonymize_pdf(path, output_path, custom_terms, field_rules, factory)
        return result_for(path, output_path, "anonymized", findings, warnings)

    if suffix in SUPPORTED_EXCEL_EXTENSIONS:
        findings, warnings = anonymize_excel(path, output_path, custom_terms, field_rules, factory)
        return result_for(path, output_path, "anonymized", findings, warnings)

    raise ValueError(f"Unsupported file extension: {suffix}")


def scan_file(
    path: Path,
    custom_terms: list[tuple[str, str]],
    field_rules: list[FieldRule],
) -> tuple[list[Finding], list[str]]:
    if path.suffix.lower() == ".docx":
        return scan_docx(path, custom_terms, field_rules)
    if path.suffix.lower() in SUPPORTED_EXCEL_EXTENSIONS:
        return scan_excel(path, custom_terms, field_rules)
    text, warnings = extract_text(path)
    return detect_sensitive_values(text, custom_terms, field_rules), warnings


def verify_file(
    path: Path,
    custom_terms: list[tuple[str, str]],
    field_rules: list[FieldRule],
) -> tuple[list[Finding], list[str]]:
    findings, warnings = scan_file(path, custom_terms, field_rules)
    return verify_findings(findings, custom_terms, field_rules), warnings


def extract_text(path: Path) -> tuple[str, list[str]]:
    suffix = path.suffix.lower()
    if suffix in SUPPORTED_TEXT_EXTENSIONS:
        return path.read_text(encoding="utf-8"), []
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX support requires python-docx. Install with: pip install python-docx") from exc
        doc = Document(path)
        parts = []
        for paragraph in doc.paragraphs:
            parts.append(paragraph.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        for paragraph in iter_docx_part_paragraphs(doc):
            parts.append(paragraph.text)
        for attribute in DOCX_METADATA_FIELDS:
            value = docx_metadata_value(doc, attribute)
            if value:
                parts.append(value)
        return "\n".join(parts), [
            "DOCX scan covers paragraphs, tables, headers, footers, and core properties; comments, text boxes, and tracked changes may need manual review.",
        ]
    if suffix == ".pdf":
        try:
            import fitz
        except ImportError as exc:
            raise RuntimeError("PDF support requires PyMuPDF. Install with: pip install pymupdf") from exc
        doc = fitz.open(path)
        text = "\n".join(page.get_text("text") for page in doc)
        metadata_text = "\n".join(
            str(doc.metadata.get(field, "") or "") for field in PDF_METADATA_FIELDS
        )
        text = "\n".join(part for part in (metadata_text, text) if part)
        warnings = []
        if not text.strip():
            warnings.append("No PDF text layer found. Use OCR or manual review for scanned pages and image-only content.")
        else:
            warnings.append("PDF scan covers extractable text and document metadata; images and annotations may need manual review.")
        return text, warnings
    raise ValueError(f"Unsupported file extension: {suffix}")


def anonymize_text(
    text: str,
    custom_terms: list[tuple[str, str]],
    field_rules: list[FieldRule],
    factory: FakeDataFactory,
) -> tuple[str, list[Finding]]:
    findings = detect_sensitive_values(text, custom_terms, field_rules)
    output = []
    cursor = 0
    for finding in non_overlapping(findings):
        output.append(text[cursor : finding.start])
        output.append(factory.replacement_for(finding.category, finding.value))
        cursor = finding.end
    output.append(text[cursor:])
    return "".join(output), findings


def verify_findings(
    findings: list[Finding],
    custom_terms: list[tuple[str, str]] | None = None,
    field_rules: list[FieldRule] | None = None,
) -> list[Finding]:
    kept: list[Finding] = []
    for finding in findings:
        if finding.source == "custom_terms":
            kept.append(finding)
            continue
        if finding.value in KNOWN_FAKE_VALUES:
            continue
        if not _still_sensitive_after_fake_removal(finding.value, custom_terms, field_rules):
            continue
        kept.append(finding)
    return kept


def _still_sensitive_after_fake_removal(
    text: str,
    custom_terms: list[tuple[str, str]] | None,
    field_rules: list[FieldRule] | None,
) -> bool:
    remainder = FAKE_VALUE_PATTERN.sub("", text)
    if remainder == text:
        return True
    stripped = remainder.strip(" \t\n、，,（）()：:;；。.-－_")
    if any(value and value in stripped for _, value in custom_terms or []):
        return True
    if not stripped:
        return False
    return any(
        finding.category in STRUCTURED_CATEGORIES
        for finding in detect_sensitive_values(stripped, custom_terms, field_rules)
    )


def non_overlapping(findings: list[Finding]) -> list[Finding]:
    selected: list[Finding] = []
    occupied_until = -1
    for finding in sorted(findings, key=lambda item: (item.start, -(item.end - item.start))):
        if finding.start >= occupied_until:
            selected.append(finding)
            occupied_until = finding.end
    return selected


NESTED_RUN_PARENTS = ("w:hyperlink", "w:sdtContent", "w:smartTag", "w:ins", "w:del", "w:fldSimple")
FIELD_MARKERS = ("w:fldChar", "w:instrText")
TEXT_CHILD_TAGS = ("w:t", "w:tab", "w:br", "w:noBreakHyphen")


def iter_paragraph_runs(paragraph) -> list:
    """Return the writable runs of a paragraph, including runs nested in hyperlinks and field results.

    ``paragraph.runs`` only exposes direct ``w:r`` children, so cached table-of-contents entries and
    hyperlink labels survive an anonymization pass that writes into direct runs only. Runs carrying
    nothing but field instructions are skipped so field codes stay intact.
    """
    try:
        from docx.oxml.ns import qn
        from docx.text.run import Run
    except ImportError:
        return list(paragraph.runs)

    elements = []
    for child in paragraph._p:
        if child.tag == qn("w:r"):
            elements.append(child)
        elif child.tag in {qn(name) for name in NESTED_RUN_PARENTS}:
            elements.extend(descendant for descendant in child.iter(qn("w:r")))
    writable = []
    for element in elements:
        if element.find(qn("w:rPr")) is None and not any(element.find(qn(tag)) is not None for tag in TEXT_CHILD_TAGS):
            continue
        if not any(element.find(qn(tag)) is not None for tag in TEXT_CHILD_TAGS):
            continue
        writable.append(element)
    return [Run(element, paragraph) for element in writable]


def paragraph_text(paragraph) -> str:
    return "".join(run.text for run in iter_paragraph_runs(paragraph))


def paragraph_has_field(paragraph) -> bool:
    try:
        from docx.oxml.ns import qn
    except ImportError:
        return False
    return any(True for _ in paragraph._p.iter(qn("w:fldChar"))) or any(
        True for _ in paragraph._p.iter(qn("w:instrText"))
    )


def anonymize_docx_paragraph(
    paragraph,
    custom_terms: list[tuple[str, str]],
    field_rules: list[FieldRule],
    label_contexts: list[str],
    factory: FakeDataFactory,
) -> tuple[list[Finding], int]:
    """Rewrite one paragraph and report how many matches could not be rewritten."""

    text = paragraph_text(paragraph)
    selected = non_overlapping(docx_paragraph_findings(text, custom_terms, field_rules, label_contexts))
    if not selected:
        return [], 0
    if paragraph_has_field(paragraph):
        # Field results such as table-of-contents entries are cached in runs that sit next to field
        # instructions. Collapsing the paragraph into one run would leave the cached original visible,
        # so rewrite each writable run in place and keep the field structure.
        for run in iter_paragraph_runs(paragraph):
            run_text = run.text
            if not run_text or not docx_paragraph_findings(run_text, custom_terms, field_rules, label_contexts):
                continue
            run_new_text, _ = anonymize_docx_paragraph_text(
                run_text, custom_terms, field_rules, label_contexts, factory
            )
            run.text = run_new_text
        remaining = paragraph_text(paragraph)
        unresolved = sum(1 for finding in selected if finding.value in remaining)
        return selected, unresolved
    new_text, _ = anonymize_docx_paragraph_text(text, custom_terms, field_rules, label_contexts, factory)
    replace_paragraph_text(paragraph, new_text)
    return selected, 0


def anonymize_docx(
    input_path: Path,
    output_path: Path,
    custom_terms: list[tuple[str, str]],
    field_rules: list[FieldRule],
    factory: FakeDataFactory,
) -> tuple[list[Finding], list[str]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("DOCX anonymization requires python-docx. Install with: pip install python-docx") from exc

    doc = Document(input_path)
    all_findings: list[Finding] = []
    warnings = [
        "DOCX anonymization covers paragraphs, tables, headers, footers, and core properties. Review comments, tracked changes, footnotes, and text boxes separately.",
        "Paragraphs with matches are rewritten as plain text runs, which may simplify styling inside those paragraphs.",
    ]

    unresolved = 0

    for paragraph in doc.paragraphs:
        findings, missed = anonymize_docx_paragraph(paragraph, custom_terms, field_rules, [], factory)
        all_findings.extend(findings)
        unresolved += missed

    for table in doc.tables:
        for row in table.rows:
            label_contexts_by_cell = adjacent_field_label_contexts(row, field_rules)
            for cell_index, cell in enumerate(row.cells):
                for paragraph in cell.paragraphs:
                    findings, missed = anonymize_docx_paragraph(
                        paragraph,
                        custom_terms,
                        field_rules,
                        label_contexts_by_cell.get(cell_index, []),
                        factory,
                    )
                    all_findings.extend(findings)
                    unresolved += missed

    for paragraph in iter_docx_part_paragraphs(doc):
        findings, missed = anonymize_docx_paragraph(paragraph, custom_terms, field_rules, [], factory)
        all_findings.extend(findings)
        unresolved += missed

    if unresolved:
        warnings.append(
            f"{unresolved} match(es) live in field or hyperlink runs that split a value across runs; "
            "they were left in place and verify mode flags them."
        )

    all_findings.extend(anonymize_docx_metadata(doc, custom_terms, field_rules, factory))

    doc.save(output_path)
    return all_findings, warnings


def anonymize_docx_metadata(doc, custom_terms, field_rules, factory) -> list[Finding]:
    findings: list[Finding] = []
    for attribute in DOCX_METADATA_FIELDS:
        value = docx_metadata_value(doc, attribute)
        if not value:
            continue
        new_value, attribute_findings = anonymize_text(value, custom_terms, field_rules, factory)
        if not attribute_findings:
            continue
        try:
            setattr(doc.core_properties, attribute, new_value)
        except (AttributeError, TypeError, ValueError):
            continue
        findings.extend(
            located_findings(
                attribute_findings,
                f"docx property {attribute}",
            )
        )
    return findings


def scan_docx(
    input_path: Path,
    custom_terms: list[tuple[str, str]],
    field_rules: list[FieldRule],
) -> tuple[list[Finding], list[str]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("DOCX scanning requires python-docx. Install with: pip install python-docx") from exc

    doc = Document(input_path)
    all_findings: list[Finding] = []
    warnings = [
        "DOCX scan covers paragraphs, tables, headers, footers, and core properties; comments, text boxes, and tracked changes may need manual review.",
    ]

    for paragraph in doc.paragraphs:
        all_findings.extend(detect_sensitive_values(paragraph_text(paragraph), custom_terms, field_rules))

    for table in doc.tables:
        for row in table.rows:
            label_contexts_by_cell = adjacent_field_label_contexts(row, field_rules)
            for cell_index, cell in enumerate(row.cells):
                for paragraph in cell.paragraphs:
                    all_findings.extend(
                        docx_paragraph_findings(
                            paragraph_text(paragraph),
                            custom_terms,
                            field_rules,
                            label_contexts_by_cell.get(cell_index, []),
                        )
                    )

    for paragraph in iter_docx_part_paragraphs(doc):
        all_findings.extend(detect_sensitive_values(paragraph_text(paragraph), custom_terms, field_rules))

    all_findings.extend(docx_metadata_findings(doc, custom_terms, field_rules))

    return all_findings, warnings


def iter_docx_part_paragraphs(doc):
    for section in doc.sections:
        for part_name in DOCX_PART_NAMES:
            part = getattr(section, part_name, None)
            if part is None:
                continue
            yield from iter_docx_story_paragraphs(part)


def iter_docx_story_paragraphs(container):
    for paragraph in getattr(container, "paragraphs", []):
        yield paragraph
    for table in getattr(container, "tables", []):
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def docx_metadata_findings(doc, custom_terms, field_rules) -> list[Finding]:
    findings: list[Finding] = []
    for attribute in DOCX_METADATA_FIELDS:
        value = docx_metadata_value(doc, attribute)
        if not value:
            continue
        findings.extend(
            located_findings(
                detect_sensitive_values(value, custom_terms, field_rules),
                f"docx property {attribute}",
            )
        )
    return findings


def docx_metadata_value(doc, attribute: str) -> str | None:
    try:
        value = getattr(doc.core_properties, attribute, None)
    except (AttributeError, ValueError):
        return None
    return str(value) if value else None


def anonymize_docx_paragraph_text(
    text: str,
    custom_terms: list[tuple[str, str]],
    field_rules: list[FieldRule],
    label_contexts: list[str],
    factory: FakeDataFactory,
) -> tuple[str, list[Finding]]:
    findings = docx_paragraph_findings(text, custom_terms, field_rules, label_contexts)
    output = []
    cursor = 0
    for finding in non_overlapping(findings):
        output.append(text[cursor : finding.start])
        output.append(factory.replacement_for(finding.category, finding.value))
        cursor = finding.end
    output.append(text[cursor:])
    return "".join(output), findings


def docx_paragraph_findings(
    text: str,
    custom_terms: list[tuple[str, str]],
    field_rules: list[FieldRule],
    label_contexts: list[str],
) -> list[Finding]:
    findings = detect_sensitive_values(text, custom_terms, field_rules)
    for label_context in label_contexts:
        for finding in detect_field_rule_values(text, field_rules, label_context=label_context):
            if not any(finding.start < item.end and finding.end > item.start for item in findings):
                findings.append(finding)
    return sorted(findings, key=lambda item: (item.start, -(item.end - item.start)))


def adjacent_field_label_contexts(row, field_rules: list[FieldRule]) -> dict[int, list[str]]:
    contexts: dict[int, list[str]] = {}
    if not field_rules:
        return contexts
    cells = row.cells
    cell_texts = [cell.text for cell in cells]
    cells_are_labels = excel_cells_are_field_labels(cell_texts, field_rules)
    for cell_index in range(len(cells) - 1):
        label_text = cell_texts[cell_index]
        value_text = cell_texts[cell_index + 1]
        if not adjacent_field_context_applies(cell_index, cell_texts, cells_are_labels):
            continue
        if not value_text.strip():
            continue
        contexts.setdefault(cell_index + 1, []).append(label_text)
    return contexts


def replace_paragraph_text(paragraph, new_text: str) -> None:
    runs = iter_paragraph_runs(paragraph)
    for run in runs:
        run.text = ""
    if runs:
        runs[0].text = new_text
    else:
        paragraph.add_run(new_text)


def anonymize_excel(
    input_path: Path,
    output_path: Path,
    custom_terms: list[tuple[str, str]],
    field_rules: list[FieldRule],
    factory: FakeDataFactory,
) -> tuple[list[Finding], list[str]]:
    workbook = load_excel_workbook(input_path)
    all_findings: list[Finding] = []
    formula_cells: list[str] = []
    sensitive_formula_cells: list[str] = []
    sensitive_unwritable_cells: list[str] = []
    sensitive_unwritable_properties: list[str] = []
    warnings = [
        "Excel anonymization covers workbook properties, worksheet titles, string cell values, hyperlink text, and cell comments (body and author), including hidden sheets.",
        "Formula expressions and non-string field-rule values are scanned but not anonymized. Review charts, images, shapes, embedded objects, pivot tables, external links, and macros separately.",
    ]
    property_findings, property_unwritable = anonymize_excel_workbook_properties(
        workbook,
        custom_terms,
        field_rules,
        factory,
    )
    all_findings.extend(property_findings)
    sensitive_unwritable_properties.extend(property_unwritable)

    for sheet_index, worksheet in enumerate(workbook.worksheets, 1):
        new_title, findings = excel_cell_anonymized_text(worksheet.title, custom_terms, field_rules, [], factory)
        if findings:
            all_findings.extend(located_findings(findings, excel_sheet_title_location(sheet_index)))
            existing_titles = [item.title for item in workbook.worksheets if item is not worksheet]
            worksheet.title = safe_excel_sheet_title(new_title, existing_titles)

    for sheet_index, worksheet in enumerate(workbook.worksheets, 1):
        vertical_label_contexts_by_column: dict[int, list[str]] = {}
        rows = list(worksheet.iter_rows())
        for row_index, row in enumerate(rows):
            if excel_row_is_blank(row):
                vertical_label_contexts_by_column = {}
                continue
            cell_texts = [excel_text_value(cell.value) for cell in row]
            cells_are_labels = excel_cells_are_field_labels(cell_texts, field_rules)
            next_row = rows[row_index + 1] if row_index + 1 < len(rows) else None
            header_contexts = excel_column_header_contexts(
                row,
                field_rules,
                cell_texts,
                cells_are_labels,
                next_row,
            )
            row_is_column_header = bool(header_contexts) and not vertical_label_contexts_by_column
            label_contexts_by_column = (
                {}
                if row_is_column_header
                else adjacent_excel_label_contexts(row, field_rules, cell_texts, cells_are_labels)
            )
            for cell_index, cell in enumerate(row):
                location = excel_cell_location(sheet_index, cell.coordinate)
                vertical_contexts = (
                    []
                    if row_is_column_header
                    or excel_cell_is_label_or_header(cell_texts[cell_index], cells_are_labels[cell_index])
                    else vertical_label_contexts_by_column.get(cell.column, [])
                )
                label_contexts = label_contexts_by_column.get(cell.column, []) + vertical_contexts
                if cell.data_type == "f":
                    if cell.value:
                        formula_cells.append(location)
                        formula_findings = located_findings(
                            excel_formula_findings(str(cell.value), custom_terms, field_rules, label_contexts),
                            location,
                        )
                        if formula_findings:
                            all_findings.extend(formula_findings)
                            sensitive_formula_cells.append(location)
                else:
                    text = excel_text_value(cell.value)
                    if text is not None:
                        new_text, findings = excel_cell_anonymized_text(
                            text,
                            custom_terms,
                            field_rules,
                            label_contexts,
                            factory,
                        )
                        if findings:
                            cell.value = new_text
                            all_findings.extend(located_findings(findings, location))
                    else:
                        detectable_text = excel_non_string_context_text(cell, label_contexts)
                        if detectable_text is not None:
                            findings = located_findings(
                                excel_cell_findings(detectable_text, custom_terms, field_rules, label_contexts),
                                location,
                            )
                            if findings:
                                all_findings.extend(findings)
                                sensitive_unwritable_cells.append(location)
                all_findings.extend(anonymize_excel_hyperlink(cell, custom_terms, field_rules, factory, location))
                all_findings.extend(anonymize_excel_comment(cell, custom_terms, field_rules, factory, location))
            if row_is_column_header and header_contexts:
                vertical_label_contexts_by_column = {
                    **vertical_label_contexts_by_column,
                    **header_contexts,
                }

    if formula_cells:
        warnings.append(
            "Formula cells were not anonymized and may contain or reveal sensitive values: "
            + ", ".join(formula_cells[:20])
            + ("..." if len(formula_cells) > 20 else "")
        )
    if sensitive_unwritable_cells:
        warnings.append(
            "Non-string Excel cells contain sensitive values that were not anonymized: "
            + ", ".join(sensitive_unwritable_cells[:20])
            + ("..." if len(sensitive_unwritable_cells) > 20 else "")
        )
    if sensitive_unwritable_properties:
        warnings.append(
            "Non-string Excel workbook properties contain sensitive values that were not anonymized: "
            + ", ".join(sensitive_unwritable_properties[:20])
            + ("..." if len(sensitive_unwritable_properties) > 20 else "")
        )
    blocking_messages = []
    if sensitive_formula_cells:
        blocking_messages.append(
            "Excel formulas contain sensitive values in cells that were not anonymized: "
            + ", ".join(sensitive_formula_cells[:20])
            + ("..." if len(sensitive_formula_cells) > 20 else "")
        )
    if sensitive_unwritable_cells:
        blocking_messages.append(
            "Excel non-string cells contain sensitive values in cells that were not anonymized: "
            + ", ".join(sensitive_unwritable_cells[:20])
            + ("..." if len(sensitive_unwritable_cells) > 20 else "")
        )
    if sensitive_unwritable_properties:
        blocking_messages.append(
            "Excel non-string workbook properties contain sensitive values that were not anonymized: "
            + ", ".join(sensitive_unwritable_properties[:20])
            + ("..." if len(sensitive_unwritable_properties) > 20 else "")
        )
    if blocking_messages:
        raise ProcessingError("; ".join(blocking_messages), all_findings, warnings)

    workbook.save(output_path)
    return all_findings, warnings


def scan_excel(
    input_path: Path,
    custom_terms: list[tuple[str, str]],
    field_rules: list[FieldRule],
) -> tuple[list[Finding], list[str]]:
    workbook = load_excel_workbook(input_path)
    all_findings: list[Finding] = []
    formula_cells: list[str] = []
    sensitive_formula_cells: list[str] = []
    sensitive_unwritable_cells: list[str] = []
    sensitive_unwritable_properties: list[str] = []
    warnings = [
        "Excel scan covers workbook properties, worksheet titles, string cell values, formula expressions, hyperlink text, and cell comments (body and author), including hidden sheets.",
        "Review non-string cells outside field-rule contexts, charts, images, shapes, embedded objects, pivot tables, external links, and macros separately.",
    ]
    property_findings, property_unwritable = scan_excel_workbook_properties(workbook, custom_terms, field_rules)
    all_findings.extend(property_findings)
    sensitive_unwritable_properties.extend(property_unwritable)

    for sheet_index, worksheet in enumerate(workbook.worksheets, 1):
        all_findings.extend(
            located_findings(
                detect_sensitive_values(worksheet.title, custom_terms, field_rules),
                excel_sheet_title_location(sheet_index),
            )
        )

    for sheet_index, worksheet in enumerate(workbook.worksheets, 1):
        vertical_label_contexts_by_column: dict[int, list[str]] = {}
        rows = list(worksheet.iter_rows())
        for row_index, row in enumerate(rows):
            if excel_row_is_blank(row):
                vertical_label_contexts_by_column = {}
                continue
            cell_texts = [excel_text_value(cell.value) for cell in row]
            cells_are_labels = excel_cells_are_field_labels(cell_texts, field_rules)
            next_row = rows[row_index + 1] if row_index + 1 < len(rows) else None
            header_contexts = excel_column_header_contexts(
                row,
                field_rules,
                cell_texts,
                cells_are_labels,
                next_row,
            )
            row_is_column_header = bool(header_contexts) and not vertical_label_contexts_by_column
            label_contexts_by_column = (
                {}
                if row_is_column_header
                else adjacent_excel_label_contexts(row, field_rules, cell_texts, cells_are_labels)
            )
            for cell_index, cell in enumerate(row):
                location = excel_cell_location(sheet_index, cell.coordinate)
                vertical_contexts = (
                    []
                    if row_is_column_header
                    or excel_cell_is_label_or_header(cell_texts[cell_index], cells_are_labels[cell_index])
                    else vertical_label_contexts_by_column.get(cell.column, [])
                )
                label_contexts = label_contexts_by_column.get(cell.column, []) + vertical_contexts
                if cell.data_type == "f":
                    if cell.value:
                        formula_cells.append(location)
                        formula_findings = located_findings(
                            excel_formula_findings(str(cell.value), custom_terms, field_rules, label_contexts),
                            location,
                        )
                        if formula_findings:
                            all_findings.extend(formula_findings)
                            sensitive_formula_cells.append(location)
                else:
                    text = excel_text_value(cell.value)
                    if text is not None:
                        all_findings.extend(
                            located_findings(
                                excel_cell_findings(
                                    text,
                                    custom_terms,
                                    field_rules,
                                    label_contexts,
                                ),
                                location,
                            )
                        )
                    else:
                        detectable_text = excel_non_string_context_text(cell, label_contexts)
                        if detectable_text is not None:
                            findings = located_findings(
                                excel_cell_findings(detectable_text, custom_terms, field_rules, label_contexts),
                                location,
                            )
                            if findings:
                                all_findings.extend(findings)
                                sensitive_unwritable_cells.append(location)
                all_findings.extend(
                    scan_excel_hyperlink(
                        cell,
                        custom_terms,
                        field_rules,
                        location,
                    )
                )
                all_findings.extend(scan_excel_comment(cell, custom_terms, field_rules, location))
            if row_is_column_header and header_contexts:
                vertical_label_contexts_by_column = {
                    **vertical_label_contexts_by_column,
                    **header_contexts,
                }

    if formula_cells:
        warnings.append(
            "Formula expressions were scanned, but calculated results were not inspected: "
            + ", ".join(formula_cells[:20])
            + ("..." if len(formula_cells) > 20 else "")
        )
    if sensitive_formula_cells:
        warnings.append(
            "Formula cells contain sensitive values and require manual review: "
            + ", ".join(sensitive_formula_cells[:20])
            + ("..." if len(sensitive_formula_cells) > 20 else "")
        )
    if sensitive_unwritable_cells:
        warnings.append(
            "Non-string cells contain sensitive values and require manual review: "
            + ", ".join(sensitive_unwritable_cells[:20])
            + ("..." if len(sensitive_unwritable_cells) > 20 else "")
        )
    if sensitive_unwritable_properties:
        warnings.append(
            "Non-string workbook properties contain sensitive values and require manual review: "
            + ", ".join(sensitive_unwritable_properties[:20])
            + ("..." if len(sensitive_unwritable_properties) > 20 else "")
        )

    return all_findings, warnings


def anonymize_excel_workbook_properties(
    workbook,
    custom_terms: list[tuple[str, str]],
    field_rules: list[FieldRule],
    factory: FakeDataFactory,
) -> tuple[list[Finding], list[str]]:
    all_findings: list[Finding] = []
    unwritable_locations: list[str] = []

    for attribute in EXCEL_WORKBOOK_PROPERTY_ATTRIBUTES:
        value = getattr(workbook.properties, attribute, None)
        location = excel_workbook_property_location(attribute)
        if isinstance(value, str):
            new_text, findings = excel_cell_anonymized_text(value, custom_terms, field_rules, [], factory)
            if findings:
                setattr(workbook.properties, attribute, new_text)
                all_findings.extend(located_findings(findings, location))
        elif value is not None:
            findings = located_findings(excel_cell_findings(str(value), custom_terms, field_rules, []), location)
            if findings:
                all_findings.extend(findings)
                unwritable_locations.append(location)

    custom_properties = getattr(workbook, "custom_doc_props", None)
    if custom_properties is not None:
        for index, prop in enumerate(list(custom_properties), 1):
            name = getattr(prop, "name", None)
            name_location = excel_custom_property_location(index, "name")
            if isinstance(name, str):
                new_name, findings = excel_cell_anonymized_text(name, custom_terms, field_rules, [], factory)
                if findings:
                    existing_names = [
                        getattr(other, "name", "")
                        for other in custom_properties
                        if other is not prop and isinstance(getattr(other, "name", None), str)
                    ]
                    prop.name = safe_excel_custom_property_name(new_name, existing_names)
                    all_findings.extend(located_findings(findings, name_location))

            value = getattr(prop, "value", None)
            value_location = excel_custom_property_location(index, "value")
            if isinstance(value, str):
                new_value, findings = excel_cell_anonymized_text(value, custom_terms, field_rules, [], factory)
                if findings:
                    prop.value = new_value
                    all_findings.extend(located_findings(findings, value_location))
            elif value is not None:
                findings = located_findings(excel_cell_findings(str(value), custom_terms, field_rules, []), value_location)
                if findings:
                    all_findings.extend(findings)
                    unwritable_locations.append(value_location)

    return all_findings, unwritable_locations


def scan_excel_workbook_properties(
    workbook,
    custom_terms: list[tuple[str, str]],
    field_rules: list[FieldRule],
) -> tuple[list[Finding], list[str]]:
    all_findings: list[Finding] = []
    unwritable_locations: list[str] = []

    for attribute in EXCEL_WORKBOOK_PROPERTY_ATTRIBUTES:
        value = getattr(workbook.properties, attribute, None)
        location = excel_workbook_property_location(attribute)
        if isinstance(value, str):
            all_findings.extend(
                located_findings(excel_cell_findings(value, custom_terms, field_rules, []), location)
            )
        elif value is not None:
            findings = located_findings(excel_cell_findings(str(value), custom_terms, field_rules, []), location)
            if findings:
                all_findings.extend(findings)
                unwritable_locations.append(location)

    custom_properties = getattr(workbook, "custom_doc_props", None)
    if custom_properties is not None:
        for index, prop in enumerate(custom_properties, 1):
            name = getattr(prop, "name", None)
            if isinstance(name, str):
                all_findings.extend(
                    located_findings(
                        excel_cell_findings(name, custom_terms, field_rules, []),
                        excel_custom_property_location(index, "name"),
                    )
                )
            value = getattr(prop, "value", None)
            value_location = excel_custom_property_location(index, "value")
            if isinstance(value, str):
                all_findings.extend(
                    located_findings(excel_cell_findings(value, custom_terms, field_rules, []), value_location)
                )
            elif value is not None:
                findings = located_findings(excel_cell_findings(str(value), custom_terms, field_rules, []), value_location)
                if findings:
                    all_findings.extend(findings)
                    unwritable_locations.append(value_location)

    return all_findings, unwritable_locations


def excel_workbook_property_location(attribute: str) -> str:
    return f"workbook property {attribute}"


def excel_custom_property_location(index: int, part: str) -> str:
    return f"workbook custom property {index} {part}"


def load_excel_workbook(path: Path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Excel support requires openpyxl. Install with: pip install openpyxl") from exc

    keep_vba = path.suffix.lower() == ".xlsm"
    return load_workbook(path, data_only=False, keep_vba=keep_vba, rich_text=True)


def excel_text_value(value) -> str | None:
    if isinstance(value, str):
        return value
    if value is not None and value.__class__.__name__ == "CellRichText":
        return str(value)
    return None


def excel_non_string_context_text(cell, label_contexts: list[str]) -> str | None:
    if not label_contexts or cell.value is None:
        return None
    if cell.data_type == "f":
        return None
    if excel_text_value(cell.value) is not None:
        return None
    return str(cell.value)


def excel_cell_location(sheet_index: int, coordinate: str) -> str:
    return f"sheet {sheet_index}!{coordinate}"


def excel_sheet_title_location(sheet_index: int) -> str:
    return f"sheet {sheet_index} title"


def located_findings(findings: list[Finding], location: str) -> list[Finding]:
    return [
        Finding(finding.category, finding.value, finding.start, finding.end, finding.source, location)
        for finding in findings
    ]


def safe_excel_sheet_title(title: str, existing_titles: list[str]) -> str:
    cleaned = INVALID_EXCEL_SHEET_TITLE_CHARS.sub("_", title).strip("'") or "Sheet"
    base = cleaned[:31] or "Sheet"
    existing = {item.lower() for item in existing_titles}
    candidate = base
    counter = 1
    while candidate.lower() in existing:
        suffix = f" {counter}"
        candidate = base[: 31 - len(suffix)] + suffix
        counter += 1
    return candidate


def safe_excel_custom_property_name(name: str, existing_names: list[str]) -> str:
    base = (name.strip() or "Property")[:255]
    existing = {item.lower() for item in existing_names}
    candidate = base
    counter = 1
    while candidate.lower() in existing:
        suffix = f" {counter}"
        candidate = base[: 255 - len(suffix)] + suffix
        counter += 1
    return candidate


def excel_cell_anonymized_text(
    text: str,
    custom_terms: list[tuple[str, str]],
    field_rules: list[FieldRule],
    label_contexts: list[str],
    factory: FakeDataFactory,
) -> tuple[str, list[Finding]]:
    findings = excel_cell_findings(text, custom_terms, field_rules, label_contexts)
    output = []
    cursor = 0
    for finding in non_overlapping(findings):
        output.append(text[cursor : finding.start])
        output.append(factory.replacement_for(finding.category, finding.value))
        cursor = finding.end
    output.append(text[cursor:])
    return "".join(output), findings


def excel_cell_findings(
    text: str,
    custom_terms: list[tuple[str, str]],
    field_rules: list[FieldRule],
    label_contexts: list[str],
) -> list[Finding]:
    findings = detect_sensitive_values(text, custom_terms, field_rules)
    for label_context in label_contexts:
        for finding in detect_field_rule_values(text, field_rules, label_context=label_context):
            if not any(finding.start < item.end and finding.end > item.start for item in findings):
                findings.append(finding)
    return sorted(findings, key=lambda item: (item.start, -(item.end - item.start)))


def excel_formula_findings(
    text: str,
    custom_terms: list[tuple[str, str]],
    field_rules: list[FieldRule],
    label_contexts: list[str],
) -> list[Finding]:
    findings = detect_sensitive_values(text, custom_terms, field_rules)
    for label_context in label_contexts:
        for literal_text, literal_start in excel_formula_string_literals(text):
            for finding in detect_field_rule_values(literal_text, field_rules, label_context=label_context):
                shifted = Finding(
                    finding.category,
                    finding.value,
                    literal_start + finding.start,
                    literal_start + finding.end,
                    "formula_field_rules",
                )
                if not any(shifted.start < item.end and shifted.end > item.start for item in findings):
                    findings.append(shifted)
    return sorted(findings, key=lambda item: (item.start, -(item.end - item.start)))


def excel_formula_string_literals(text: str) -> list[tuple[str, int]]:
    literals: list[tuple[str, int]] = []
    cursor = 0
    while cursor < len(text):
        if text[cursor] != '"':
            cursor += 1
            continue
        literal_start = cursor + 1
        cursor = literal_start
        while cursor < len(text):
            if text[cursor] == '"':
                if cursor + 1 < len(text) and text[cursor + 1] == '"':
                    cursor += 2
                    continue
                literals.append((text[literal_start:cursor], literal_start))
                cursor += 1
                break
            cursor += 1
    return literals


def adjacent_excel_label_contexts(
    row,
    field_rules: list[FieldRule],
    cell_texts: list[str | None] | None = None,
    cells_are_labels: list[bool] | None = None,
) -> dict[int, list[str]]:
    contexts: dict[int, list[str]] = {}
    if not field_rules:
        return contexts
    cell_texts = cell_texts if cell_texts is not None else [excel_text_value(cell.value) for cell in row]
    cells_are_labels = cells_are_labels if cells_are_labels is not None else excel_cells_are_field_labels(cell_texts, field_rules)
    for cell_index in range(len(row) - 1):
        label = cell_texts[cell_index]
        value_cell = row[cell_index + 1]
        value_text = cell_texts[cell_index + 1]
        if label is None or not adjacent_field_context_applies(cell_index, cell_texts, cells_are_labels):
            continue
        if value_cell.value is None:
            continue
        if value_text is not None and looks_like_excel_table_header_value(value_text):
            continue
        contexts.setdefault(value_cell.column, []).append(label)
    return contexts


def excel_cells_are_field_labels(cell_texts: list[str | None], field_rules: list[FieldRule]) -> list[bool]:
    if not field_rules:
        return [False for _text in cell_texts]
    return [
        bool(text is not None and any(field_rule_matches_label(rule, text) for rule in field_rules))
        for text in cell_texts
    ]


def excel_row_is_blank(row) -> bool:
    return all(cell.value is None and cell.hyperlink is None for cell in row)


def excel_column_header_contexts(
    row,
    field_rules: list[FieldRule],
    cell_texts: list[str | None],
    cells_are_labels: list[bool],
    next_row=None,
) -> dict[int, list[str]]:
    if not field_rules:
        return {}
    nonblank_indexes = [
        index
        for index, text in enumerate(cell_texts)
        if text is not None and normalize_excel_label(text)
    ]
    label_indexes = [index for index in nonblank_indexes if cells_are_labels[index]]
    if not nonblank_indexes or not label_indexes:
        return {}

    row_looks_like_header = (
        (len(nonblank_indexes) == 1 and excel_row_has_only_value(row, nonblank_indexes[0]))
        or (
            len(nonblank_indexes) > 1
            and all(cells_are_labels[index] or excel_cell_is_header_value(cell_texts[index]) for index in nonblank_indexes)
        )
        or excel_next_row_supports_column_headers(
            row,
            label_indexes,
            field_rules,
            next_row,
        )
    )
    if not row_looks_like_header:
        return {}
    return {
        cell.column: [cell_texts[index]]
        for index, cell in enumerate(row)
        if cells_are_labels[index] and cell_texts[index] is not None
    }


def excel_row_has_only_value(row, value_index: int) -> bool:
    return all(index == value_index or (cell.value is None and cell.hyperlink is None) for index, cell in enumerate(row))


def excel_next_row_supports_column_headers(
    row,
    label_indexes: list[int],
    field_rules: list[FieldRule],
    next_row,
) -> bool:
    if next_row is None:
        return False
    for index in label_indexes:
        if index >= len(next_row):
            continue
        label_text = excel_text_value(row[index].value)
        value_text = excel_context_probe_text(next_row[index])
        if label_text is None or value_text is None:
            continue
        if any(field_rule_matches_label(rule, value_text) for rule in field_rules):
            continue
        if excel_cell_is_header_value(value_text):
            continue
        if excel_field_rule_context_findings(value_text, field_rules, label_text):
            return True
    return False


def excel_field_rule_context_findings(text: str, field_rules: list[FieldRule], label_context: str) -> list[Finding]:
    findings: list[Finding] = []
    for finding in detect_field_rule_values(text, field_rules, label_context=label_context):
        if not any(finding.start < item.end and finding.end > item.start for item in findings):
            findings.append(finding)
    return sorted(findings, key=lambda item: (item.start, -(item.end - item.start)))


def excel_context_probe_text(cell) -> str | None:
    text = excel_text_value(cell.value)
    if text is not None:
        return text
    if cell.value is None or cell.data_type == "f":
        return None
    return str(cell.value)


def excel_cell_is_label_or_header(text: str | None, is_label: bool) -> bool:
    return is_label or excel_cell_is_header_value(text)


def excel_cell_is_header_value(text: str | None) -> bool:
    return bool(text is not None and looks_like_excel_table_header_value(text))


def looks_like_excel_table_header_value(text: str) -> bool:
    return normalize_excel_label(text) in EXCEL_TABLE_HEADER_VALUES


def normalize_excel_label(text: str) -> str:
    return re.sub(r"[\s:：；;，,。]", "", text or "")


def adjacent_field_context_applies(
    cell_index: int,
    cell_texts: list[str | None],
    cells_are_labels: list[bool],
) -> bool:
    if cell_index >= len(cell_texts) - 1:
        return False
    label = cell_texts[cell_index]
    value_text = cell_texts[cell_index + 1]
    if label is None or not cells_are_labels[cell_index]:
        return False
    if cell_index > 0 and cells_are_labels[cell_index - 1]:
        return False
    if cells_are_labels[cell_index + 1]:
        return False
    return not (value_text is not None and looks_like_excel_table_header_value(value_text))


def anonymize_excel_hyperlink(
    cell,
    custom_terms: list[tuple[str, str]],
    field_rules: list[FieldRule],
    factory: FakeDataFactory,
    cell_location: str,
) -> list[Finding]:
    hyperlink = cell.hyperlink
    if hyperlink is None:
        return []

    findings: list[Finding] = []
    for attribute in EXCEL_HYPERLINK_TEXT_ATTRIBUTES:
        value = getattr(hyperlink, attribute, None)
        if not isinstance(value, str):
            continue
        new_text, attribute_findings = excel_cell_anonymized_text(value, custom_terms, field_rules, [], factory)
        if attribute_findings:
            setattr(hyperlink, attribute, new_text)
            findings.extend(located_findings(attribute_findings, f"{cell_location} hyperlink {attribute}"))
    return findings


def excel_comment_parts(cell) -> list[tuple[str, str]]:
    """Return (attribute, text) pairs for the comment attached to a cell, if any."""

    comment = getattr(cell, "comment", None)
    if comment is None:
        return []
    return [
        (attribute, str(value))
        for attribute, value in (("author", comment.author), ("text", comment.text))
        if value
    ]


def anonymize_excel_comment(
    cell,
    custom_terms: list[tuple[str, str]],
    field_rules: list[FieldRule],
    factory: FakeDataFactory,
    cell_location: str,
) -> list[Finding]:
    """Rewrite comment bodies and comment authors, which openpyxl keeps as separate text."""

    comment = getattr(cell, "comment", None)
    if comment is None:
        return []
    findings: list[Finding] = []
    for attribute, value in excel_comment_parts(cell):
        new_text, attribute_findings = excel_cell_anonymized_text(value, custom_terms, field_rules, [], factory)
        if not attribute_findings:
            continue
        setattr(comment, attribute, new_text)
        findings.extend(located_findings(attribute_findings, f"{cell_location} comment {attribute}"))
    return findings


def scan_excel_comment(
    cell,
    custom_terms: list[tuple[str, str]],
    field_rules: list[FieldRule],
    cell_location: str,
) -> list[Finding]:
    findings: list[Finding] = []
    for attribute, value in excel_comment_parts(cell):
        findings.extend(
            located_findings(
                excel_cell_findings(value, custom_terms, field_rules, []),
                f"{cell_location} comment {attribute}",
            )
        )
    return findings


def scan_excel_hyperlink(
    cell,
    custom_terms: list[tuple[str, str]],
    field_rules: list[FieldRule],
    cell_location: str,
) -> list[Finding]:
    hyperlink = cell.hyperlink
    if hyperlink is None:
        return []

    findings: list[Finding] = []
    for attribute in EXCEL_HYPERLINK_TEXT_ATTRIBUTES:
        value = getattr(hyperlink, attribute, None)
        if isinstance(value, str):
            findings.extend(
                located_findings(
                    excel_cell_findings(value, custom_terms, field_rules, []),
                    f"{cell_location} hyperlink {attribute}",
                )
            )
    return findings


def anonymize_pdf(
    input_path: Path,
    output_path: Path,
    custom_terms: list[tuple[str, str]],
    field_rules: list[FieldRule],
    factory: FakeDataFactory,
) -> tuple[list[Finding], list[str]]:
    try:
        import fitz
    except ImportError as exc:
        raise RuntimeError("PDF anonymization requires PyMuPDF. Install with: pip install pymupdf") from exc

    doc = fitz.open(input_path)
    all_findings: list[Finding] = []
    warnings = [
        "PDF anonymization uses redaction annotations and inserts fake text near matched rectangles.",
        "PDF document metadata is rewritten when it contains detected values.",
        "Review scanned pages, images, annotations, forms, and unusual layouts manually.",
    ]

    for page in doc:
        text = page.get_text("text")
        if not text.strip():
            warnings.append(f"Page {page.number + 1} has no extractable text layer.")
            continue
        findings = non_overlapping(detect_sensitive_values(text, custom_terms, field_rules))
        all_findings.extend(findings)
        annotations_added = False
        missing_redactions = []
        searched_values = set()
        for finding in findings:
            if finding.value in searched_values:
                continue
            searched_values.add(finding.value)
            replacement = factory.replacement_for(finding.category, finding.value)
            rects = list(page.search_for(finding.value))
            if not rects:
                missing_redactions.append(
                    f"page {page.number + 1} {normalize_category(finding.category)} length {len(finding.value)}"
                )
                continue
            for rect in rects:
                # CJK 替换值必须显式指定中文字体：PyMuPDF 默认 Helvetica 没有中文字形，
                # 否则替换文本会渲染成 "?"（敏感值仍被删除，但产物不可读、不可检索）。
                page.add_redact_annot(
                    rect,
                    text=replacement,
                    fill=(1, 1, 1),
                    text_color=(0, 0, 0),
                    fontname="china-s",
                )
                annotations_added = True
        if missing_redactions:
            raise RuntimeError(
                "PDF redaction could not locate rectangles for detected text: " + "; ".join(missing_redactions)
            )
        if annotations_added:
            page.apply_redactions()

    all_findings.extend(anonymize_pdf_metadata(doc, custom_terms, field_rules, factory))

    doc.save(output_path, garbage=4, deflate=True)
    return all_findings, warnings


def anonymize_pdf_metadata(doc, custom_terms, field_rules, factory) -> list[Finding]:
    findings: list[Finding] = []
    metadata = dict(doc.metadata or {})
    updated = False
    for field in PDF_METADATA_FIELDS:
        value = metadata.get(field)
        if not value:
            continue
        new_value, field_findings = anonymize_text(str(value), custom_terms, field_rules, factory)
        if not field_findings:
            continue
        metadata[field] = new_value
        updated = True
        findings.extend(located_findings(field_findings, f"pdf property {field}"))
    if updated:
        doc.set_metadata(metadata)
    return findings


def plan_outputs(
    inputs: list[Path],
    args: argparse.Namespace,
    custom_terms: list[tuple[str, str]] | None = None,
    field_rules: list[FieldRule] | None = None,
    factory: FakeDataFactory | None = None,
) -> dict[Path, Path]:
    if args.mode != "anonymize":
        return {}

    outputs: dict[Path, Path] = {}
    seen: dict[Path, Path] = {}
    roots = input_roots(args.inputs)
    for path in inputs:
        output_path = output_for(path, args, roots, len(inputs), custom_terms, field_rules, factory)
        validate_output_extension(path, output_path)
        if same_path(path, output_path):
            raise ValueError(f"Refusing to overwrite original file: {safe_report_path(path)}")
        resolved_output = output_path.expanduser().resolve(strict=False)
        if resolved_output in seen:
            if not args.anonymize_filenames:
                raise ValueError(
                    "Output path collision: "
                    + f"{safe_report_path(seen[resolved_output])} and {safe_report_path(path)} "
                    + f"both map to {safe_report_path(output_path)}"
                )
            output_path = deduplicated_output(output_path, path)
            resolved_output = output_path.expanduser().resolve(strict=False)
        seen[resolved_output] = path
        outputs[path] = output_path
    return outputs


def deduplicated_output(output_path: Path, input_path: Path) -> Path:
    digest = hashlib.sha256(str(input_path.expanduser().resolve(strict=False)).encode("utf-8")).hexdigest()[:8]
    return output_path.with_name(f"{output_path.stem}-{digest}{output_path.suffix}")


def output_for(
    path: Path,
    args: argparse.Namespace,
    roots: list[Path],
    input_count: int,
    custom_terms: list[tuple[str, str]] | None = None,
    field_rules: list[FieldRule] | None = None,
    factory: FakeDataFactory | None = None,
) -> Path:
    if args.output:
        if input_count > 1:
            raise ValueError("--output can only be used with one input file.")
        return args.output
    if not args.anonymize_filenames:
        terms, rules, name_factory = None, None, None
    else:
        terms, rules, name_factory = custom_terms, field_rules, factory
    if args.output_dir:
        return args.output_dir / relative_output_name(path, roots, terms, rules, name_factory)
    return path.with_name(anonymized_name(path, terms, rules, name_factory))


def validate_output_extension(input_path: Path, output_path: Path) -> None:
    input_suffix = input_path.suffix.lower()
    output_suffix = output_path.suffix.lower()
    if input_suffix in SUPPORTED_EXCEL_EXTENSIONS and output_suffix != input_suffix:
        raise ValueError(
            "Excel output extension must match input extension: "
            + f"{safe_report_path(input_path)} -> {safe_report_path(output_path)}"
        )


def anonymized_name(
    path: Path,
    custom_terms: list[tuple[str, str]] | None = None,
    field_rules: list[FieldRule] | None = None,
    factory: FakeDataFactory | None = None,
) -> str:
    stem = path.stem
    if factory is not None:
        stem, _ = anonymize_text(stem, custom_terms or [], field_rules or [], factory)
    return f"{stem}.anonymized{path.suffix}"


def relative_output_name(
    path: Path,
    roots: list[Path] | None,
    custom_terms: list[tuple[str, str]] | None = None,
    field_rules: list[FieldRule] | None = None,
    factory: FakeDataFactory | None = None,
) -> Path:
    resolved_path = path.expanduser().resolve(strict=False)
    for root in roots or []:
        try:
            relative = resolved_path.relative_to(root)
        except ValueError:
            continue
        return relative.with_name(anonymized_name(relative, custom_terms, field_rules, factory))
    return Path(anonymized_name(path, custom_terms, field_rules, factory))


def input_roots(inputs: list[Path]) -> list[Path]:
    roots = []
    for item in inputs:
        if item.is_dir():
            roots.append(item.expanduser().resolve(strict=False))
    return sorted(roots, key=lambda path: len(path.parts), reverse=True)


def same_path(left: Path, right: Path) -> bool:
    left_expanded = left.expanduser()
    right_expanded = right.expanduser()
    if left_expanded.resolve(strict=False) == right_expanded.resolve(strict=False):
        return True
    try:
        return left_expanded.samefile(right_expanded)
    except FileNotFoundError:
        return False
    except OSError:
        return False


def result_for(
    input_path: Path,
    output_path: Path | None,
    status: str,
    findings: list[Finding],
    warnings: list[str],
) -> dict:
    counts: dict[str, int] = {}
    for finding in findings:
        category = normalize_category(finding.category)
        counts[category] = counts.get(category, 0) + 1

    result = {
        "input_path": safe_report_path(input_path),
        "input_extension": input_path.suffix.lower(),
        "status": status,
        "finding_counts": counts,
        "finding_count": len(findings),
        "findings": [
            {
                "category": normalize_category(finding.category),
                "source": finding.source,
                "start": finding.start,
                "end": finding.end,
                "length": len(finding.value),
                **({"location": finding.location} if finding.location else {}),
            }
            for finding in findings
        ],
        "warnings": warnings,
    }
    if output_path:
        result["output_path"] = safe_report_path(output_path)
        result["output_extension"] = output_path.suffix.lower()
    return result


def safe_report_path(path: Path) -> str:
    resolved = str(path.expanduser().resolve(strict=False))
    digest = hashlib.sha256(resolved.encode("utf-8")).hexdigest()[:16]
    suffix = path.suffix.lower()
    return f"path:{digest}{suffix}"


def safe_error_message(message: str, *paths: Path) -> str:
    sanitized = message
    for path in paths:
        safe_path = safe_report_path(path)
        candidates = {
            str(path),
            str(path.expanduser()),
            str(path.expanduser().resolve(strict=False)),
        }
        for candidate in candidates:
            sanitized = sanitized.replace(candidate, safe_path)
    return sanitized


def build_report(
    results: list[dict],
    factory: FakeDataFactory,
    skipped_inputs: list[dict] | None = None,
) -> dict:
    del factory
    skipped_inputs = skipped_inputs or []
    summary_counts: dict[str, int] = {}
    for result in results:
        for category, count in result.get("finding_counts", {}).items():
            summary_counts[category] = summary_counts.get(category, 0) + count

    report = {
        "summary": {
            "files_processed": len(results),
            "total_findings": sum(result.get("finding_count", 0) for result in results),
            "finding_counts": summary_counts,
            "errors": sum(1 for result in results if result.get("status") == "error"),
            **skipped_summary(skipped_inputs),
        },
        "outputs": results,
    }
    if skipped_inputs:
        report["skipped_inputs"] = skipped_inputs
    return report


def skipped_summary(skipped_inputs: list[dict]) -> dict:
    return {
        "skipped_inputs": len(skipped_inputs),
        "skipped_unsupported": sum(1 for item in skipped_inputs if item.get("reason") == "unsupported_extension"),
        "skipped_missing": sum(1 for item in skipped_inputs if item.get("reason") == "missing"),
    }


def exit_code_for(results: list[dict], skipped_inputs: list[dict]) -> int:
    if any(result.get("status") == "error" for result in results):
        return 1
    if any(result.get("status") == "residual_findings" for result in results):
        return 1
    if any(item.get("explicit") for item in skipped_inputs):
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
