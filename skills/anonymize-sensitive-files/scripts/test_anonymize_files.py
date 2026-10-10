import json
import subprocess
import sys
import tempfile
import types
import unittest
import zipfile
from copy import copy
from pathlib import Path
from unittest.mock import patch

import anonymize_files

SCRIPT = Path(__file__).with_name("anonymize_files.py")


class AnonymizeFilesTest(unittest.TestCase):
    def test_markdown_anonymization_uses_fake_data_without_leaking_originals(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "case.md"
            source.write_text(
                "\n".join(
                    [
                        "# 客户资料",
                        "姓名：李雷",
                        "联系人：李雷",
                        "手机号：13812345678",
                        "邮箱：lilei@example.com",
                        "身份证号：110105199001010002",
                        "API Key: sk-live-abcdefghijklmnopqrstuvwxyz123456",
                    ]
                ),
                encoding="utf-8",
            )

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--seed",
                    "20260603",
                    "--report",
                    str(tmp_path / "report.json"),
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            payload = json.loads(result.stdout)
            output = tmp_path / "case.anonymized.md"
            report = json.loads((tmp_path / "report.json").read_text(encoding="utf-8"))
            anonymized = output.read_text(encoding="utf-8")

            self.assertTrue(output.exists())
            self.assertEqual(source.read_text(encoding="utf-8").count("李雷"), 2)
            self.assertNotIn("李雷", anonymized)
            self.assertNotIn("13812345678", anonymized)
            self.assertNotIn("lilei@example.com", anonymized)
            self.assertNotIn("110105199001010002", anonymized)
            self.assertNotIn("sk-live-abcdefghijklmnopqrstuvwxyz123456", anonymized)
            self.assertNotIn("[REDACTED", anonymized)
            self.assertIn("19999999999", anonymized)
            self.assertEqual(anonymized.count("张三"), 2)
            self.assertEqual(report["summary"]["total_findings"], 6)
            self.assertNotIn("李雷", json.dumps(report, ensure_ascii=False))
            self.assertNotIn("13812345678", json.dumps(report, ensure_ascii=False))

    def test_scan_mode_does_not_create_an_output_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "note.txt"
            source.write_text("电话：13812345678\n", encoding="utf-8")

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--mode",
                    "scan",
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            payload = json.loads(result.stdout)
            self.assertEqual(payload["summary"]["files_processed"], 1)
            self.assertEqual(payload["summary"]["total_findings"], 1)
            self.assertFalse((tmp_path / "note.anonymized.txt").exists())

    def test_custom_terms_are_replaced_with_stable_fake_values(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "memo.md"
            terms = tmp_path / "terms.txt"
            source.write_text("客户：星河集团\n项目：天枢计划\n星河集团复核。\n", encoding="utf-8")
            terms.write_text("org:星河集团\nproject:天枢计划\n", encoding="utf-8")

            subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--terms",
                    str(terms),
                    "--seed",
                    "1",
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            anonymized = (tmp_path / "memo.anonymized.md").read_text(encoding="utf-8")
            self.assertNotIn("星河集团", anonymized)
            self.assertNotIn("天枢计划", anonymized)
            self.assertEqual(anonymized.count("北京星河科技有限公司"), 2)
            self.assertIn("星云迁移项目", anonymized)

    def test_global_field_rules_anonymize_labeled_text_without_terms_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "fields.md"
            source.write_text(
                "\n".join(
                    [
                        "财务联系人：顾安琪",
                        "联行号：313305007918",
                        "收件人：沈亦然",
                    ]
                ),
                encoding="utf-8",
            )

            subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--seed",
                    "field-rules",
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            anonymized = (tmp_path / "fields.anonymized.md").read_text(encoding="utf-8")
            self.assertNotIn("顾安琪", anonymized)
            self.assertNotIn("313305007918", anonymized)
            self.assertNotIn("沈亦然", anonymized)

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(tmp_path / "fields.anonymized.md"),
                    "--mode",
                    "verify",
                ],
                text=True,
                capture_output=True,
                check=True,
            )
            payload = json.loads(result.stdout)
            self.assertEqual(payload["summary"]["total_findings"], 0)

    def test_global_field_rules_anonymize_labeled_mobile_with_separators(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "phone.md"
            source.write_text("财务手机：138 1234 5678\n", encoding="utf-8")

            subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            anonymized = (tmp_path / "phone.anonymized.md").read_text(encoding="utf-8")
            self.assertNotIn("138 1234 5678", anonymized)
            self.assertIn("19999999999", anonymized)

    def test_global_field_rules_anonymize_markdown_table_adjacent_cells(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "table-fields.md"
            source.write_text(
                "\n".join(
                    [
                        "| 字段 | 值 | 字段 | 值 |",
                        "| --- | --- | --- | --- |",
                        "| 法定代表人 | 林若航 | 采购联系人 | 顾安琪 |",
                        "| 紧急联系人 | 沈亦然 138-1234-5678 | 销售联系人 | 唐若云 |",
                    ]
                ),
                encoding="utf-8",
            )

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--seed",
                    "markdown-table-field-rules",
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            payload = json.loads(result.stdout)
            anonymized = (tmp_path / "table-fields.anonymized.md").read_text(encoding="utf-8")
            for original in ["林若航", "顾安琪", "沈亦然", "唐若云", "138-1234-5678"]:
                self.assertNotIn(original, anonymized)
            self.assertIn("19999999999", anonymized)
            self.assertNotRegex(anonymized, r"1[3-9]\d-\d{4}-\d{4}")
            self.assertEqual(payload["summary"]["finding_counts"], {"name": 4, "phone": 1})

    def test_verify_reports_residual_markdown_table_adjacent_cells(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "residual-table.md"
            source.write_text(
                "\n".join(
                    [
                        "| 字段 | 值 |",
                        "| --- | --- |",
                        "| 法定代表人 | 林若航 |",
                    ]
                ),
                encoding="utf-8",
            )

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--mode",
                    "verify",
                ],
                text=True,
                capture_output=True,
                check=False,
            )

            payload = json.loads(result.stdout)
            self.assertEqual(result.returncode, 1)
            self.assertEqual(payload["summary"]["finding_counts"], {"name": 1})

    def test_markdown_table_adjacent_rules_ignore_sentence_cells(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "sentence-table.md"
            source.write_text(
                "\n".join(
                    [
                        "| 条款 | 条件 | 说明 |",
                        "| --- | --- | --- |",
                        "| 付款条件 | 合同生效并收到乙方合格发票 | 甲方信息 |",
                        "| 甲方 | 乙方 | 签署栏 | 盖章 |",
                        "| 项目 | 甲方 | 项目 | 乙方 |",
                    ]
                ),
                encoding="utf-8",
            )

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--mode",
                    "scan",
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            payload = json.loads(result.stdout)
            self.assertEqual(payload["summary"]["total_findings"], 0)

    def test_markdown_table_adjacent_phone_label_does_not_match_contact_name_rule(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "contact-phone-table.md"
            source.write_text(
                "\n".join(
                    [
                        "| 字段 | 值 |",
                        "| --- | --- |",
                        "| 联系人手机号 | 010-12345678 |",
                    ]
                ),
                encoding="utf-8",
            )

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            payload = json.loads(result.stdout)
            anonymized = (tmp_path / "contact-phone-table.anonymized.md").read_text(encoding="utf-8")
            self.assertEqual(payload["summary"]["finding_counts"], {"phone": 1})
            self.assertIn("19999999999", anonymized)
            self.assertNotIn("张三", anonymized)

    def test_expanded_global_field_rules_cover_common_business_fields(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "expanded.md"
            source.write_text(
                "\n".join(
                    [
                        "法定代表人：林若航",
                        "甲方：北京北辰采购有限公司",
                        "开票邮箱：invoice@example.com",
                        "纳税人识别号：91110000000000000X",
                        "运单号：SF1234567890",
                        "微信号：wxid_demo_12345",
                        "MAC地址：AA:BB:CC:DD:EE:FF",
                    ]
                ),
                encoding="utf-8",
            )

            subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--seed",
                    "expanded-field-rules",
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            anonymized = (tmp_path / "expanded.anonymized.md").read_text(encoding="utf-8")
            for original in [
                "林若航",
                "北京北辰采购有限公司",
                "invoice@example.com",
                "91110000000000000X",
                "SF1234567890",
                "wxid_demo_12345",
                "AA:BB:CC:DD:EE:FF",
            ]:
                self.assertNotIn(original, anonymized)
            self.assertIn("纳税人识别号：", anonymized)
            self.assertIn("MAC地址：", anonymized)

    def test_global_field_rules_anonymize_docx_adjacent_table_cells(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx is not installed")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "fields.docx"
            doc = Document()
            table = doc.add_table(rows=3, cols=2)
            table.cell(0, 0).text = "收件人"
            table.cell(0, 1).text = "沈亦然"
            table.cell(1, 0).text = "财务联系人"
            table.cell(1, 1).text = "顾安琪"
            table.cell(2, 0).text = "联行号"
            table.cell(2, 1).text = "313305007918"
            doc.save(source)

            subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--seed",
                    "field-rules",
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            output = tmp_path / "fields.anonymized.docx"
            anonymized_doc = Document(output)
            text = "\n".join(
                cell.text
                for table in anonymized_doc.tables
                for row in table.rows
                for cell in row.cells
            )
            self.assertNotIn("沈亦然", text)
            self.assertNotIn("顾安琪", text)
            self.assertNotIn("313305007918", text)

    def test_docx_adjacent_rules_ignore_neighboring_label_cells(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx is not installed")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "label-row.docx"
            doc = Document()
            table = doc.add_table(rows=1, cols=4)
            for index, value in enumerate(["甲方", "乙方", "签署栏", "盖章"]):
                table.cell(0, index).text = value
            doc.save(source)

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            payload = json.loads(result.stdout)
            output = tmp_path / "label-row.anonymized.docx"
            anonymized_doc = Document(output)
            cells = [anonymized_doc.tables[0].cell(0, index).text for index in range(4)]

            self.assertEqual(payload["summary"]["total_findings"], 0)
            self.assertEqual(cells, ["甲方", "乙方", "签署栏", "盖章"])

    def test_docx_adjacent_phone_label_uses_phone_replacement_for_landline(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx is not installed")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "phone-table.docx"
            doc = Document()
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "联系人电话"
            table.cell(0, 1).text = "010-12345678"
            doc.save(source)

            subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            output = tmp_path / "phone-table.anonymized.docx"
            anonymized_doc = Document(output)
            self.assertEqual(anonymized_doc.tables[0].cell(0, 1).text, "19999999999")

    def test_xlsx_anonymizes_cells_and_adjacent_field_values(self):
        try:
            from openpyxl import Workbook, load_workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "customers.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "客户"
            sheet["A1"] = "姓名"
            sheet["B1"] = "李雷"
            sheet["A2"] = "手机号"
            sheet["B2"] = "13812345678"
            sheet["A3"] = "邮箱：lilei@example.com"
            sheet["A4"] = "=SUM(1, 2)"
            sheet["A5"] = 13812345678
            font = copy(sheet["A1"].font)
            font.bold = True
            sheet["A1"].font = font
            workbook.create_sheet("隐藏")
            workbook["隐藏"]["A1"] = "联系人：李雷"
            workbook["隐藏"].sheet_state = "hidden"
            workbook.save(source)

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--seed",
                    "xlsx",
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            payload = json.loads(result.stdout)
            output = tmp_path / "customers.anonymized.xlsx"
            anonymized = load_workbook(output, data_only=False)
            sheet = anonymized["客户"]
            hidden = anonymized["隐藏"]

            self.assertTrue(output.exists())
            self.assertEqual(payload["summary"]["finding_counts"], {"name": 2, "phone": 1, "email": 1})
            self.assertEqual(sheet["A1"].value, "姓名")
            self.assertEqual(sheet["B1"].value, "张三")
            self.assertEqual(sheet["B2"].value, "19999999999")
            self.assertNotIn("lilei@example.com", sheet["A3"].value)
            self.assertEqual(sheet["A4"].value, "=SUM(1, 2)")
            self.assertEqual(sheet["A5"].value, 13812345678)
            self.assertTrue(sheet["A1"].font.bold)
            self.assertEqual(hidden.sheet_state, "hidden")
            self.assertEqual(hidden["A1"].value, "联系人：张三")
            self.assertTrue(any("Formula cells" in warning for warning in payload["outputs"][0]["warnings"]))

            verify = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(output),
                    "--mode",
                    "verify",
                ],
                text=True,
                capture_output=True,
                check=True,
            )
            self.assertEqual(json.loads(verify.stdout)["outputs"][0]["status"], "verified")

    def test_xlsx_mixed_column_headers_anonymize_values_without_rewriting_headers(self):
        try:
            from openpyxl import Workbook, load_workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "mixed-headers.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["姓名", "部门", "手机号"])
            sheet.append(["李雷", "销售", "13812345678"])
            workbook.save(source)

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            payload = json.loads(result.stdout)
            output = tmp_path / "mixed-headers.anonymized.xlsx"
            anonymized = load_workbook(output)
            rows = list(anonymized.active.iter_rows(values_only=True))
            locations = {finding["location"] for finding in payload["outputs"][0]["findings"]}

            self.assertEqual(payload["summary"]["finding_counts"], {"name": 1, "phone": 1})
            self.assertEqual(rows[0], ("姓名", "部门", "手机号"))
            self.assertEqual(rows[1], ("张三", "销售", "19999999999"))
            self.assertIn("sheet 1!A2", locations)
            self.assertIn("sheet 1!C2", locations)
            self.assertNotIn("sheet 1!B1", locations)

            verify = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(output),
                    "--mode",
                    "verify",
                ],
                text=True,
                capture_output=True,
                check=True,
            )
            self.assertEqual(json.loads(verify.stdout)["outputs"][0]["status"], "verified")

    def test_xlsx_adjacent_numeric_values_are_reported_and_block_anonymize(self):
        try:
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "numeric-phone.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = "手机号"
            sheet["B1"] = 13812345678
            workbook.save(source)

            scan = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--mode",
                    "scan",
                ],
                text=True,
                capture_output=True,
                check=True,
            )
            self.assertEqual(json.loads(scan.stdout)["summary"]["finding_counts"], {"phone": 1})

            anonymize = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                ],
                text=True,
                capture_output=True,
                check=False,
            )

            payload = json.loads(anonymize.stdout)
            combined_output = anonymize.stdout + anonymize.stderr
            self.assertNotEqual(anonymize.returncode, 0)
            self.assertEqual(payload["outputs"][0]["status"], "error")
            self.assertEqual(payload["summary"]["finding_counts"], {"phone": 1})
            self.assertFalse((tmp_path / "numeric-phone.anonymized.xlsx").exists())
            self.assertNotIn("13812345678", combined_output)
            self.assertIn("sheet 1!B1", combined_output)

            verify = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--mode",
                    "verify",
                ],
                text=True,
                capture_output=True,
                check=False,
            )
            verify_payload = json.loads(verify.stdout)
            self.assertNotEqual(verify.returncode, 0)
            self.assertEqual(verify_payload["outputs"][0]["status"], "residual_findings")
            self.assertEqual(verify_payload["summary"]["finding_counts"], {"phone": 1})

    def test_xlsx_adjacent_formula_string_values_are_reported_and_block_anonymize(self):
        try:
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "formula-name.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = "收件人"
            sheet["B1"] = '="沈亦然"'
            workbook.save(source)

            scan = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--mode",
                    "scan",
                ],
                text=True,
                capture_output=True,
                check=True,
            )
            self.assertEqual(json.loads(scan.stdout)["summary"]["finding_counts"], {"name": 1})

            anonymize = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                ],
                text=True,
                capture_output=True,
                check=False,
            )

            payload = json.loads(anonymize.stdout)
            combined_output = anonymize.stdout + anonymize.stderr
            self.assertNotEqual(anonymize.returncode, 0)
            self.assertEqual(payload["outputs"][0]["status"], "error")
            self.assertEqual(payload["summary"]["finding_counts"], {"name": 1})
            self.assertFalse((tmp_path / "formula-name.anonymized.xlsx").exists())
            self.assertNotIn("沈亦然", combined_output)
            self.assertIn("sheet 1!B1", combined_output)

            verify = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--mode",
                    "verify",
                ],
                text=True,
                capture_output=True,
                check=False,
            )
            verify_payload = json.loads(verify.stdout)
            self.assertNotEqual(verify.returncode, 0)
            self.assertEqual(verify_payload["outputs"][0]["status"], "residual_findings")
            self.assertEqual(verify_payload["summary"]["finding_counts"], {"name": 1})

    def test_xlsx_anonymizes_sheet_titles_without_leaking_them_in_warnings(self):
        try:
            from openpyxl import Workbook, load_workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "sheet-title.xlsx"
            terms = tmp_path / "terms.txt"
            terms.write_text("name:李雷\n", encoding="utf-8")
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "李雷13812345678"
            sheet["A1"] = "=SUM(1, 2)"
            sheet["A2"] = "姓名"
            sheet["B2"] = "李雷"
            workbook.save(source)

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--terms",
                    str(terms),
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            payload = json.loads(result.stdout)
            report_text = json.dumps(payload, ensure_ascii=False)
            output = tmp_path / "sheet-title.anonymized.xlsx"
            anonymized = load_workbook(output, data_only=False)

            self.assertNotIn("李雷", report_text)
            self.assertNotIn("13812345678", report_text)
            self.assertNotIn("李雷", "".join(anonymized.sheetnames))
            self.assertNotIn("13812345678", "".join(anonymized.sheetnames))
            self.assertTrue(any("sheet 1!A1" in warning for warning in payload["outputs"][0]["warnings"]))

    def test_xlsx_anonymize_fails_without_output_when_formula_contains_sensitive_value(self):
        try:
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "formula.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = '=CONCAT("李雷", "13812345678")'
            workbook.save(source)

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                ],
                text=True,
                capture_output=True,
                check=False,
            )

            payload = json.loads(result.stdout)
            output = tmp_path / "formula.anonymized.xlsx"
            combined_output = result.stdout + result.stderr
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(payload["outputs"][0]["status"], "error")
            self.assertEqual(payload["summary"]["finding_counts"], {"phone": 1})
            self.assertFalse(output.exists())
            self.assertNotIn("李雷", combined_output)
            self.assertNotIn("13812345678", combined_output)
            self.assertIn("sheet 1!A1", combined_output)

    def test_xlsx_verify_reports_residual_sensitive_formula(self):
        try:
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "formula-residual.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = '=CONCAT("李雷", "13812345678")'
            workbook.save(source)

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--mode",
                    "verify",
                ],
                text=True,
                capture_output=True,
                check=False,
            )

            payload = json.loads(result.stdout)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(payload["outputs"][0]["status"], "residual_findings")
            self.assertEqual(payload["summary"]["finding_counts"], {"phone": 1})
            self.assertNotIn("李雷", result.stdout)
            self.assertNotIn("13812345678", result.stdout)

    def test_xlsx_anonymizes_hyperlink_targets_without_raw_residue(self):
        try:
            from openpyxl import Workbook, load_workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "hyperlink.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = "邮箱：lilei@example.com"
            sheet["A1"].hyperlink = "mailto:lilei@example.com"
            sheet["A1"].style = "Hyperlink"
            workbook.save(source)

            subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            output = tmp_path / "hyperlink.anonymized.xlsx"
            anonymized = load_workbook(output)
            target = anonymized.active["A1"].hyperlink.target
            with zipfile.ZipFile(output) as package:
                package_text = "\n".join(
                    package.read(name).decode("utf-8", errors="ignore") for name in package.namelist()
                )

            self.assertNotIn("lilei@example.com", anonymized.active["A1"].value)
            self.assertNotIn("lilei@example.com", target)
            self.assertNotIn("lilei@example.com", package_text)

            verify = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(output),
                    "--mode",
                    "verify",
                ],
                text=True,
                capture_output=True,
                check=True,
            )
            self.assertEqual(json.loads(verify.stdout)["outputs"][0]["status"], "verified")

    def test_xlsx_adjacent_rules_ignore_neighboring_label_cells(self):
        try:
            from openpyxl import Workbook, load_workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "labels.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["甲方", "乙方", "签署栏", "盖章"])
            sheet.append(["字段", "值", "字段", "值"])
            workbook.save(source)

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            payload = json.loads(result.stdout)
            output = tmp_path / "labels.anonymized.xlsx"
            anonymized = load_workbook(output)
            rows = list(anonymized.active.iter_rows(values_only=True))

            self.assertEqual(payload["summary"]["total_findings"], 0)
            self.assertEqual(rows[0], ("甲方", "乙方", "签署栏", "盖章"))
            self.assertEqual(rows[1], ("字段", "值", "字段", "值"))

    def test_xlsx_column_header_values_are_scanned_and_anonymized(self):
        try:
            from openpyxl import Workbook, load_workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "column-name.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = "姓名"
            sheet["A2"] = "李雷"
            workbook.save(source)

            scan = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--mode",
                    "scan",
                ],
                text=True,
                capture_output=True,
                check=True,
            )
            self.assertEqual(json.loads(scan.stdout)["summary"]["finding_counts"], {"name": 1})

            subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            output = tmp_path / "column-name.anonymized.xlsx"
            anonymized = load_workbook(output)
            self.assertEqual(anonymized.active["A1"].value, "姓名")
            self.assertEqual(anonymized.active["A2"].value, "张三")

    def test_xlsx_column_header_numeric_values_are_reported_and_block_anonymize(self):
        try:
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "column-phone.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = "手机号"
            sheet["A2"] = 13812345678
            workbook.save(source)

            scan = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--mode",
                    "scan",
                ],
                text=True,
                capture_output=True,
                check=True,
            )
            self.assertEqual(json.loads(scan.stdout)["summary"]["finding_counts"], {"phone": 1})

            anonymize = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                ],
                text=True,
                capture_output=True,
                check=False,
            )

            payload = json.loads(anonymize.stdout)
            combined_output = anonymize.stdout + anonymize.stderr
            self.assertNotEqual(anonymize.returncode, 0)
            self.assertEqual(payload["outputs"][0]["status"], "error")
            self.assertEqual(payload["summary"]["finding_counts"], {"phone": 1})
            self.assertFalse((tmp_path / "column-phone.anonymized.xlsx").exists())
            self.assertNotIn("13812345678", combined_output)
            self.assertIn("sheet 1!A2", combined_output)

    def test_xlsx_anonymizes_workbook_properties_and_verify_scans_them(self):
        try:
            from openpyxl import Workbook, load_workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "metadata.xlsx"
            terms = tmp_path / "terms.txt"
            terms.write_text("name:李雷\norg:星河集团\n", encoding="utf-8")
            workbook = Workbook()
            workbook.properties.creator = "李雷"
            workbook.properties.title = "星河集团客户资料 13812345678"
            workbook.active["A1"] = "普通内容"
            workbook.save(source)

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--terms",
                    str(terms),
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            payload = json.loads(result.stdout)
            output = tmp_path / "metadata.anonymized.xlsx"
            anonymized = load_workbook(output)
            properties_text = "\n".join(
                str(value or "")
                for value in [
                    anonymized.properties.creator,
                    anonymized.properties.title,
                    anonymized.properties.subject,
                    anonymized.properties.description,
                    anonymized.properties.keywords,
                    anonymized.properties.category,
                    anonymized.properties.contentStatus,
                    anonymized.properties.lastModifiedBy,
                ]
            )

            self.assertEqual(payload["summary"]["finding_counts"], {"name": 1, "phone": 1, "org": 1})
            self.assertNotIn("李雷", properties_text)
            self.assertNotIn("星河集团", properties_text)
            self.assertNotIn("13812345678", properties_text)

            verify = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(output),
                    "--mode",
                    "verify",
                    "--terms",
                    str(terms),
                ],
                text=True,
                capture_output=True,
                check=True,
            )
            self.assertEqual(json.loads(verify.stdout)["outputs"][0]["status"], "verified")

    def test_xlsx_preserves_rich_text_cells_without_findings(self):
        try:
            from openpyxl import Workbook, load_workbook
            from openpyxl.cell.rich_text import CellRichText, TextBlock
            from openpyxl.cell.text import InlineFont
        except ImportError:
            self.skipTest("openpyxl rich text support is not installed")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "rich-text.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = CellRichText(TextBlock(InlineFont(b=True), "普通"), "说明")
            workbook.save(source)

            subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            output = tmp_path / "rich-text.anonymized.xlsx"
            anonymized = load_workbook(output, rich_text=True)
            value = anonymized.active["A1"].value

            self.assertIsInstance(value, CellRichText)
            self.assertIsInstance(value[0], TextBlock)
            self.assertTrue(value[0].font.b)

    def test_xlsm_preserves_vba_project_part(self):
        try:
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")

        def add_vba_project(workbook_path: Path) -> None:
            with zipfile.ZipFile(workbook_path, "r") as package:
                entries = {
                    item.filename: package.read(item.filename)
                    for item in package.infolist()
                    if item.filename not in {"[Content_Types].xml", "xl/_rels/workbook.xml.rels", "xl/vbaProject.bin"}
                }
                content_types = package.read("[Content_Types].xml").decode("utf-8")
                if "vbaProject.bin" not in content_types:
                    content_types = content_types.replace(
                        "</Types>",
                        '<Override PartName="/xl/vbaProject.bin" '
                        'ContentType="application/vnd.ms-office.vbaProject"/></Types>',
                    )
                rels = package.read("xl/_rels/workbook.xml.rels").decode("utf-8")
                if "vbaProject" not in rels:
                    rels = rels.replace(
                        "</Relationships>",
                        '<Relationship Id="rIdVBA" '
                        'Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" '
                        'Target="vbaProject.bin"/></Relationships>',
                    )
            with zipfile.ZipFile(workbook_path, "w") as package:
                for name, content in entries.items():
                    package.writestr(name, content)
                package.writestr("[Content_Types].xml", content_types)
                package.writestr("xl/_rels/workbook.xml.rels", rels)
                package.writestr("xl/vbaProject.bin", b"fake-vba-project")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "macro.xlsm"
            workbook = Workbook()
            workbook.active["A1"] = "电话：13812345678"
            workbook.save(source)
            add_vba_project(source)

            subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            output = tmp_path / "macro.anonymized.xlsm"
            with zipfile.ZipFile(output) as package:
                self.assertEqual(package.read("xl/vbaProject.bin"), b"fake-vba-project")
                content_types = package.read("[Content_Types].xml").decode("utf-8")
                rels = package.read("xl/_rels/workbook.xml.rels").decode("utf-8")

            self.assertIn("vbaProject", content_types)
            self.assertIn("macroEnabled", content_types)
            self.assertIn("vbaProject", rels)

    def test_xlsm_output_extension_must_remain_xlsm(self):
        try:
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "macro.xlsm"
            output = tmp_path / "macro.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "电话：13812345678"
            workbook.save(source)

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--output",
                    str(output),
                ],
                text=True,
                capture_output=True,
                check=False,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertFalse(output.exists())
            self.assertIn("Excel output extension", result.stderr + result.stdout)

    def test_xls_is_reported_as_unsupported(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "legacy.xls"
            source.write_bytes(b"not a real workbook")

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--mode",
                    "scan",
                ],
                text=True,
                capture_output=True,
                check=False,
            )

            payload = json.loads(result.stderr)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(payload["skipped_inputs"][0]["reason"], "unsupported_extension")
            self.assertEqual(payload["skipped_inputs"][0]["extension"], ".xls")

    def test_report_paths_are_sanitized_by_default(self):
        try:
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "李雷-客户资料.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "普通内容"
            workbook.save(source)

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--mode",
                    "scan",
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            payload = json.loads(result.stdout)
            report_text = json.dumps(payload, ensure_ascii=False)
            self.assertNotIn("李雷", report_text)
            self.assertNotIn("客户资料", report_text)
            self.assertNotEqual(payload["outputs"][0]["input_path"], str(source))
            self.assertEqual(payload["outputs"][0]["input_extension"], ".xlsx")

    def test_mapping_file_hashes_originals_by_default(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "note.txt"
            mapping = tmp_path / "mapping.json"
            source.write_text("联系人：李雷\n电话：13812345678\n", encoding="utf-8")

            subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--mapping",
                    str(mapping),
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            mapping_text = mapping.read_text(encoding="utf-8")
            self.assertNotIn("李雷", mapping_text)
            self.assertNotIn("13812345678", mapping_text)
            self.assertIn("张三", mapping_text)

    def test_output_cannot_overwrite_original_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "李雷-note.txt"
            original = "电话：13812345678\n"
            source.write_text(original, encoding="utf-8")

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--output",
                    str(source),
                ],
                text=True,
                capture_output=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(source.read_text(encoding="utf-8"), original)
            self.assertIn("overwrite", result.stderr + result.stdout)
            self.assertNotIn("李雷", result.stderr + result.stdout)

    def test_output_cannot_overwrite_original_file_through_hardlink(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "source.txt"
            output = tmp_path / "output.txt"
            original = "电话：13812345678\n"
            source.write_text(original, encoding="utf-8")
            output.hardlink_to(source)

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--output",
                    str(output),
                ],
                text=True,
                capture_output=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(source.read_text(encoding="utf-8"), original)
            self.assertEqual(output.read_text(encoding="utf-8"), original)
            self.assertIn("overwrite", result.stderr + result.stdout)

    def test_output_path_is_sanitized_in_runtime_errors(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "note.txt"
            output = tmp_path / "李雷-output.txt"
            source.write_text("电话：13812345678\n", encoding="utf-8")
            output.mkdir()

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--output",
                    str(output),
                ],
                text=True,
                capture_output=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertNotIn("李雷-output", result.stdout + result.stderr)
            self.assertIn("path:", result.stdout + result.stderr)

    def test_terms_and_field_rules_path_errors_are_sanitized_json(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "note.txt"
            source.write_text("电话：13812345678\n", encoding="utf-8")

            terms_result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--terms",
                    str(tmp_path / "李雷-terms.txt"),
                    "--mode",
                    "scan",
                ],
                text=True,
                capture_output=True,
            )
            field_rules_result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--field-rules",
                    str(tmp_path / "星河-rules.yaml"),
                    "--mode",
                    "scan",
                ],
                text=True,
                capture_output=True,
            )

            for result in (terms_result, field_rules_result):
                self.assertNotEqual(result.returncode, 0)
                payload = json.loads(result.stderr)
                self.assertIn("error", payload)
                self.assertNotIn("李雷", result.stdout + result.stderr)
                self.assertNotIn("星河", result.stdout + result.stderr)
                self.assertIn("path:", result.stdout + result.stderr)

    def test_output_dir_preserves_relative_paths_for_duplicate_names(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source_a = tmp_path / "north" / "case.txt"
            source_b = tmp_path / "south" / "case.txt"
            output_dir = tmp_path / "out"
            source_a.parent.mkdir()
            source_b.parent.mkdir()
            source_a.write_text("电话：13812345678\n", encoding="utf-8")
            source_b.write_text("电话：13912345678\n", encoding="utf-8")

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(tmp_path),
                    "--recursive",
                    "--output-dir",
                    str(output_dir),
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            json.loads(result.stdout)
            self.assertTrue((output_dir / "north" / "case.anonymized.txt").exists())
            self.assertTrue((output_dir / "south" / "case.anonymized.txt").exists())
            self.assertIn("19999999999", (output_dir / "north" / "case.anonymized.txt").read_text(encoding="utf-8"))
            self.assertIn("18888888888", (output_dir / "south" / "case.anonymized.txt").read_text(encoding="utf-8"))

    def test_verify_ignores_values_generated_by_anonymizer(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "note.txt"
            source.write_text("电话：13812345678\n邮箱：lilei@example.com\n", encoding="utf-8")
            subprocess.run(
                [sys.executable, str(SCRIPT), str(source)],
                text=True,
                capture_output=True,
                check=True,
            )

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(tmp_path / "note.anonymized.txt"),
                    "--mode",
                    "verify",
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            payload = json.loads(result.stdout)
            self.assertEqual(payload["summary"]["total_findings"], 0)
            self.assertEqual(payload["outputs"][0]["status"], "verified")

    def test_verify_ignores_generated_secret_replacements(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "secret.txt"
            original_secret = "sk-live-abcdefghijklmnopqrstuvwxyz123456"
            source.write_text(f"API Key: {original_secret}\n", encoding="utf-8")

            subprocess.run(
                [sys.executable, str(SCRIPT), str(source)],
                text=True,
                capture_output=True,
                check=True,
            )

            anonymized_path = tmp_path / "secret.anonymized.txt"
            anonymized = anonymized_path.read_text(encoding="utf-8")
            self.assertNotIn(original_secret, anonymized)

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(anonymized_path),
                    "--mode",
                    "verify",
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            payload = json.loads(result.stdout)
            self.assertEqual(payload["summary"]["total_findings"], 0)
            self.assertEqual(payload["outputs"][0]["status"], "verified")

    def test_verify_returns_nonzero_when_residual_findings_found(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "note.txt"
            terms = tmp_path / "terms.txt"
            source.write_text("客户：星河集团\n", encoding="utf-8")
            terms.write_text("org:星河集团\n", encoding="utf-8")

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--mode",
                    "verify",
                    "--terms",
                    str(terms),
                ],
                text=True,
                capture_output=True,
            )

            payload = json.loads(result.stdout)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(payload["outputs"][0]["status"], "residual_findings")
            self.assertEqual(payload["summary"]["total_findings"], 1)

    def test_export_mapping_does_not_put_raw_values_in_stdout_or_report(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "note.txt"
            mapping = tmp_path / "mapping.json"
            report = tmp_path / "report.json"
            source.write_text("联系人：李雷\n电话：13812345678\n", encoding="utf-8")

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--mapping",
                    str(mapping),
                    "--export-mapping",
                    "--report",
                    str(report),
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            report_text = report.read_text(encoding="utf-8")
            mapping_text = mapping.read_text(encoding="utf-8")
            self.assertNotIn("李雷", result.stdout)
            self.assertNotIn("13812345678", result.stdout)
            self.assertNotIn("李雷", report_text)
            self.assertNotIn("13812345678", report_text)
            self.assertIn("李雷", mapping_text)
            self.assertIn("13812345678", mapping_text)

    def test_export_mapping_requires_mapping_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "note.txt"
            source.write_text("联系人：李雷\n", encoding="utf-8")

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--export-mapping",
                ],
                text=True,
                capture_output=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("--export-mapping requires --mapping", result.stderr)
            self.assertNotIn("李雷", result.stdout + result.stderr)

    def test_recursive_directory_reports_unsupported_files(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            supported = tmp_path / "note.txt"
            unsupported = tmp_path / "legacy.doc"
            supported.write_text("电话：13812345678\n", encoding="utf-8")
            unsupported.write_text("secret", encoding="utf-8")

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(tmp_path),
                    "--recursive",
                    "--mode",
                    "scan",
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            payload = json.loads(result.stdout)
            self.assertEqual(payload["summary"]["skipped_unsupported"], 1)
            self.assertNotEqual(payload["skipped_inputs"][0]["path"], str(unsupported))
            self.assertEqual(payload["skipped_inputs"][0]["extension"], ".doc")
            self.assertEqual(payload["skipped_inputs"][0]["reason"], "unsupported_extension")

    def test_explicit_unsupported_input_returns_nonzero_and_reports_skip(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            unsupported = tmp_path / "legacy.doc"
            unsupported.write_text("secret", encoding="utf-8")

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(unsupported),
                    "--mode",
                    "scan",
                ],
                text=True,
                capture_output=True,
            )

            payload = json.loads(result.stderr)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(payload["summary"]["skipped_unsupported"], 1)
            self.assertEqual(payload["skipped_inputs"][0]["explicit"], True)

    def test_pdf_anonymization_errors_when_detected_text_has_no_redaction_rectangles(self):
        class FakePage:
            number = 0

            def get_text(self, mode):
                if mode != "text":
                    raise AssertionError(f"unexpected text extraction mode: {mode}")
                return "电话：13812345678"

            def search_for(self, value):
                if value != "13812345678":
                    raise AssertionError(f"unexpected search value: {value}")
                return []

            def add_redact_annot(self, *args, **kwargs):
                raise AssertionError("redaction annotation should not be added without rectangles")

            def apply_redactions(self):
                raise AssertionError("redactions should not be applied without rectangles")

        class FakeDoc:
            def __init__(self):
                self.page = FakePage()
                self.saved = False

            def __iter__(self):
                return iter([self.page])

            def save(self, *args, **kwargs):
                self.saved = True

        fake_doc = FakeDoc()
        fake_fitz = types.SimpleNamespace(open=lambda path: fake_doc)

        with patch.dict(sys.modules, {"fitz": fake_fitz}):
            with self.assertRaisesRegex(RuntimeError, "could not locate"):
                anonymize_files.anonymize_pdf(
                    Path("input.pdf"),
                    Path("output.pdf"),
                    [],
                    [],
                    anonymize_files.FakeDataFactory(),
                )
        self.assertFalse(fake_doc.saved)

    def test_custom_terms_do_not_keep_original_when_original_matches_first_fake_value(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "memo.md"
            terms = tmp_path / "terms.txt"
            source.write_text("公司：北京星河科技有限公司\n项目：星云迁移项目\n", encoding="utf-8")
            terms.write_text("org:北京星河科技有限公司\nproject:星云迁移项目\n", encoding="utf-8")

            subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--terms",
                    str(terms),
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            anonymized = (tmp_path / "memo.anonymized.md").read_text(encoding="utf-8")
            self.assertNotIn("北京星河科技有限公司", anonymized)
            self.assertNotIn("星云迁移项目", anonymized)


    def test_docx_anonymizes_headers_and_core_properties(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx is not installed")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "memo.docx"
            terms = tmp_path / "terms.txt"
            doc = Document()
            doc.core_properties.title = "星河集团档案"
            doc.sections[0].header.paragraphs[0].text = "星河集团内部资料"
            doc.add_paragraph("联系人：李雷")
            doc.save(source)
            terms.write_text("org:星河集团\n", encoding="utf-8")

            subprocess.run(
                [sys.executable, str(SCRIPT), str(source), "--terms", str(terms)],
                text=True,
                capture_output=True,
                check=True,
            )

            anonymized_doc = Document(tmp_path / "memo.anonymized.docx")
            self.assertNotIn("星河集团", anonymized_doc.core_properties.title)
            self.assertNotIn("星河集团", anonymized_doc.sections[0].header.paragraphs[0].text)
            self.assertIn("张三", anonymized_doc.paragraphs[0].text)

    def test_pdf_anonymizes_document_metadata(self):
        try:
            import fitz
        except ImportError:
            self.skipTest("PyMuPDF is not installed")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "brief.pdf"
            terms = tmp_path / "terms.txt"
            doc = fitz.open()
            page = doc.new_page()
            page.insert_text((72, 72), "Quarterly brief")
            doc.set_metadata({"title": "星河集团并购简报", "author": "星河集团"})
            doc.save(source)
            doc.close()
            terms.write_text("org:星河集团\n", encoding="utf-8")

            subprocess.run(
                [sys.executable, str(SCRIPT), str(source), "--terms", str(terms)],
                text=True,
                capture_output=True,
                check=True,
            )

            output = fitz.open(tmp_path / "brief.anonymized.pdf")
            metadata_text = json.dumps(output.metadata, ensure_ascii=False)
            output.close()
            self.assertNotIn("星河集团", metadata_text)

    def test_anonymize_filenames_replaces_entities_in_output_file_name(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "FS-02_启阳项目-社保公积金统计表.md"
            terms = tmp_path / "terms.txt"
            source.write_text("项目名称：启阳项目\n", encoding="utf-8")
            terms.write_text("project:启阳项目\n", encoding="utf-8")

            subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(source),
                    "--terms",
                    str(terms),
                    "--anonymize-filenames",
                ],
                text=True,
                capture_output=True,
                check=True,
            )

            outputs = sorted(path.name for path in tmp_path.glob("*.anonymized.md"))
            self.assertEqual(len(outputs), 1)
            self.assertNotIn("启阳", outputs[0])
            self.assertTrue(outputs[0].startswith("FS-02_"))

    def test_file_names_are_kept_when_anonymize_filenames_is_not_used(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "FS-02_启阳项目-社保公积金统计表.md"
            terms = tmp_path / "terms.txt"
            source.write_text("项目名称：启阳项目\n", encoding="utf-8")
            terms.write_text("project:启阳项目\n", encoding="utf-8")

            subprocess.run(
                [sys.executable, str(SCRIPT), str(source), "--terms", str(terms)],
                text=True,
                capture_output=True,
                check=True,
            )

            outputs = sorted(path.name for path in tmp_path.glob("*.anonymized.md"))
            self.assertEqual(outputs, ["FS-02_启阳项目-社保公积金统计表.anonymized.md"])


    def test_replacement_does_not_echo_the_original_brand_through_fake_pools(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "brand.md"
            terms = tmp_path / "terms.txt"
            source.write_text("甲方：星河集团\n乙方：星河集团复核\n", encoding="utf-8")
            terms.write_text("org:星河集团\norg:星河\n", encoding="utf-8")

            subprocess.run(
                [sys.executable, str(SCRIPT), str(source), "--terms", str(terms)],
                text=True,
                capture_output=True,
                check=True,
            )

            anonymized = (tmp_path / "brand.anonymized.md").read_text(encoding="utf-8")
            self.assertNotIn("星河", anonymized)
            self.assertIn("甲方：", anonymized)


    def test_xlsx_column_header_context_survives_label_like_data_cells(self):
        try:
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "ledger.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = "项目名称"
            sheet["B1"] = "律师事务所"
            sheet["A2"] = "曲江观邸"
            sheet["B2"] = "陕西锦路律师事务所"
            sheet["A3"] = "馨苑"
            sheet["B3"] = "陕西锦路律师事务所"
            workbook.save(source)
            terms = tmp_path / "terms.txt"
            terms.write_text("org:陕西锦路律师事务所\n", encoding="utf-8")

            subprocess.run(
                [sys.executable, str(SCRIPT), str(source), "--terms", str(terms)],
                text=True,
                capture_output=True,
                check=True,
            )

            from openpyxl import load_workbook

            anonymized = load_workbook(tmp_path / "ledger.anonymized.xlsx").active
            self.assertNotEqual(anonymized["A2"].value, "曲江观邸")
            self.assertNotEqual(anonymized["A3"].value, "馨苑")
            self.assertNotIn("锦路", str(anonymized["B2"].value) + str(anonymized["B3"].value))


    def test_verify_treats_fake_value_plus_non_sensitive_suffix_as_clean(self):
        finding = anonymize_files.Finding("name", "张敏（男性）", 0, 6, "field_rules", "联系人")
        self.assertEqual(anonymize_files.verify_findings([finding], [], []), [])

        leftover = anonymize_files.Finding("name", "某某某公司", 0, 5, "field_rules", "联系人")
        self.assertEqual(len(anonymize_files.verify_findings([leftover], [], [])), 1)

    def test_traversal_order_does_not_change_fake_value_assignment(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source_dir = tmp_path / "corpus"
            source_dir.mkdir()
            (source_dir / "a.md").write_text("项目名称：启阳项目\n承租人：融辉集团\n", encoding="utf-8")
            (source_dir / "b.md").write_text("法定代表人：张三\n", encoding="utf-8")
            (source_dir / "c.md").write_text("启阳项目与融辉集团的补充协议\n", encoding="utf-8")
            terms = tmp_path / "terms.txt"
            terms.write_text("project:启阳项目\norg:融辉集团\nname:张三\n", encoding="utf-8")
            recursive_map = tmp_path / "map-recursive.json"
            explicit_map = tmp_path / "map-explicit.json"

            subprocess.run(
                [
                    sys.executable, str(SCRIPT), str(source_dir), "--recursive",
                    "--terms", str(terms),
                    "--output-dir", str(tmp_path / "out-recursive"),
                    "--mapping", str(recursive_map), "--export-mapping",
                ],
                text=True, capture_output=True, check=True,
            )
            subprocess.run(
                [
                    sys.executable, str(SCRIPT),
                    str(source_dir / "c.md"), str(source_dir / "b.md"), str(source_dir / "a.md"),
                    "--terms", str(terms),
                    "--output-dir", str(tmp_path / "out-explicit"),
                    "--mapping", str(explicit_map), "--export-mapping",
                ],
                text=True, capture_output=True, check=True,
            )

            self.assertEqual(
                json.loads(recursive_map.read_text(encoding="utf-8")),
                json.loads(explicit_map.read_text(encoding="utf-8")),
            )
            for name in ("a.md", "b.md", "c.md"):
                first = (tmp_path / "out-recursive" / f"{Path(name).stem}.anonymized.md").read_text(encoding="utf-8")
                second = (tmp_path / "out-explicit" / f"{Path(name).stem}.anonymized.md").read_text(encoding="utf-8")
                self.assertEqual(first, second, name)

    def test_explicit_duplicate_input_paths_are_processed_once(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "note.md"
            source.write_text("联系人：张三\n", encoding="utf-8")

            result = subprocess.run(
                [
                    sys.executable, str(SCRIPT), str(source), str(source),
                    "--output-dir", str(tmp_path / "out"), "--report", str(tmp_path / "report.json"),
                ],
                text=True, capture_output=True, check=True,
            )

            report = json.loads((tmp_path / "report.json").read_text(encoding="utf-8"))
            self.assertEqual(report["summary"]["files_processed"], 1)
            self.assertEqual(len(list((tmp_path / "out").iterdir())), 1)

    def test_xlsx_anonymizes_cell_comment_body_and_author(self):
        from openpyxl import Workbook, load_workbook
        from openpyxl.comments import Comment

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "ledger.xlsx"
            terms = tmp_path / "terms.txt"
            book = Workbook()
            sheet = book.active
            sheet["A1"] = "项目名称"
            sheet["B1"] = "启阳项目"
            sheet["B1"].comment = Comment("启阳项目由刘秀军负责", "刘秀军")
            book.save(source)
            terms.write_text("project:启阳项目\nname:刘秀军\n", encoding="utf-8")

            subprocess.run(
                [
                    sys.executable, str(SCRIPT), str(source),
                    "--terms", str(terms), "--output-dir", str(tmp_path / "out"),
                    "--report", str(tmp_path / "report.json"),
                ],
                text=True, capture_output=True, check=True,
            )

            output = next((tmp_path / "out").glob("*.anonymized.xlsx"))
            comment = load_workbook(output).active["B1"].comment
            rendered = f"{comment.author} {comment.text}"
            self.assertNotIn("启阳项目", rendered)
            self.assertNotIn("刘秀军", rendered)
            self.assertNotIn("启阳项目", str(load_workbook(output).active["B1"].value))
            locations = [
                finding["location"]
                for finding in json.loads((tmp_path / "report.json").read_text(encoding="utf-8"))["outputs"][0]["findings"]
            ]
            self.assertTrue(any(item.endswith("comment author") for item in locations), locations)
            self.assertTrue(any(item.endswith("comment text") for item in locations), locations)

    def test_docx_field_cached_hyperlink_text_is_anonymized_without_breaking_the_field(self):
        from docx import Document
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "opinion.docx"
            terms = tmp_path / "terms.txt"

            document = Document()
            document.add_paragraph("正文：法定代表人张三。")
            paragraph = document.add_paragraph()
            for chunk in (
                f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>',
                f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve"> TOC \\o "1-3" </w:instrText></w:r>',
                f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="separate"/></w:r>',
                f'<w:hyperlink {nsdecls("w")}>'
                "<w:r><w:t>关于</w:t></w:r>"
                "<w:r><w:t>金辉控股</w:t></w:r>"
                "<w:r><w:t>的说明</w:t></w:r>"
                "</w:hyperlink>",
                f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>',
            ):
                paragraph._p.append(parse_xml(chunk))
            document.save(source)
            terms.write_text("org:金辉控股\nname:张三\n", encoding="utf-8")

            subprocess.run(
                [
                    sys.executable, str(SCRIPT), str(source),
                    "--terms", str(terms), "--output-dir", str(tmp_path / "out"),
                ],
                text=True, capture_output=True, check=True,
            )

            output = next((tmp_path / "out").glob("*.anonymized.docx"))
            with zipfile.ZipFile(output) as archive:
                xml = archive.read("word/document.xml").decode("utf-8")
            self.assertNotIn("金辉控股", xml)
            self.assertNotIn("张三", xml)
            self.assertIn('w:instrText', xml)
            self.assertIn("TOC", xml)
            self.assertIn("关于", xml)
            self.assertIn("的说明", xml)



if __name__ == "__main__":
    unittest.main()
